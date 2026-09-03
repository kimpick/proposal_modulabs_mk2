#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
AI 캠퍼스 훈련운영계획서 Word 템플릿 생성기
- aicampus_proposal.pdf 구조를 기반으로 한 재사용 가능한 Word 템플릿
"""

import os
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import copy

# ──────────────────────────────────────────
# 색상 정의 (PDF 기반)
# ──────────────────────────────────────────
PRIMARY_BLUE = RGBColor(0x20, 0x3A, 0x7A)    # 장 제목 배경, 표 헤더
ACCENT_BLUE = RGBColor(0x1F, 0x4E, 0x79)     # 강조 텍스트
DARK_GRAY = RGBColor(0x40, 0x40, 0x40)       # 본문 텍스트
MED_GRAY = RGBColor(0x7F, 0x7F, 0x7F)        # 부가 정보
LIGHT_GRAY = RGBColor(0x80, 0x80, 0x80)      # 설명 텍스트
HEADER_BG = "203A7A"                           # 표 헤더 배경 (hex)
SUBHEADER_BG = "D6E4F0"                       # 표 서브헤더 배경
LIGHT_BG = "F2F2F2"                           # 표 짝수행 배경
RED_ACCENT = RGBColor(0xC0, 0x00, 0x00)       # 강조 (빨강)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# ──────────────────────────────────────────
# 폰트 헬퍼
# ──────────────────────────────────────────
FONT_KR = "맑은 고딕"      # Malgun Gothic
FONT_KR_BOLD = "맑은 고딕"
FONT_TITLE = "맑은 고딕"    # 원본은 HYgtrE (고딕계열)

def set_run_font(run, size_pt, bold=False, color=DARK_GRAY, font_name=FONT_KR):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    # 한글 폰트 설정
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)


def add_paragraph_styled(doc, text, size_pt, bold=False, color=DARK_GRAY,
                         alignment=WD_ALIGN_PARAGRAPH.LEFT,
                         space_before=0, space_after=0,
                         font_name=FONT_KR):
    p = doc.add_paragraph()
    p.alignment = alignment
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = Pt(size_pt * 1.6)
    run = p.add_run(text)
    set_run_font(run, size_pt, bold=bold, color=color, font_name=font_name)
    return p


def set_cell_shading(cell, color_hex):
    """표 셀 배경색 설정"""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_text(cell, text, size_pt=10, bold=False, color=DARK_GRAY,
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, bg_color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = alignment
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    pf.line_spacing = Pt(size_pt * 1.5)
    run = p.add_run(text)
    set_run_font(run, size_pt, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if bg_color:
        set_cell_shading(cell, bg_color)


def create_styled_table(doc, headers, rows, col_widths=None):
    """스타일이 적용된 테이블 생성"""
    num_cols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    # 헤더 행
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, header, size_pt=10, bold=True, color=WHITE, bg_color=HEADER_BG)

    # 데이터 행
    for r_idx, row_data in enumerate(rows):
        for c_idx, cell_text in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            bg = LIGHT_BG if r_idx % 2 == 1 else None
            align = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cell, str(cell_text), size_pt=9, alignment=align, bg_color=bg)

    # 열 너비 설정
    if col_widths:
        for row in table.rows:
            for i, width in enumerate(col_widths):
                row.cells[i].width = Cm(width)

    return table


def add_page_break(doc):
    doc.add_page_break()


# ──────────────────────────────────────────
# 메인 템플릿 생성
# ──────────────────────────────────────────
def generate_template():
    doc = Document()

    # ── 페이지 설정 ──
    section = doc.sections[0]
    section.page_width = Cm(21.0)       # A4
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # ════════════════════════════════════════
    # 표지 (Cover Page)
    # ════════════════════════════════════════
    # 상단 여백
    for _ in range(6):
        doc.add_paragraph()

    # AI 캠퍼스 로고/배지 영역
    add_paragraph_styled(doc, "AI 캠퍼스", 14, bold=True, color=WHITE,
                         alignment=WD_ALIGN_PARAGRAPH.LEFT)
    # 배지 배경은 표로 대체
    badge_table = doc.add_table(rows=1, cols=1)
    badge_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    badge_cell = badge_table.rows[0].cells[0]
    set_cell_shading(badge_cell, HEADER_BG)
    set_cell_text(badge_cell, "AI 캠퍼스", size_pt=14, bold=True, color=WHITE,
                  alignment=WD_ALIGN_PARAGRAPH.CENTER)
    badge_cell.width = Cm(4)

    doc.add_paragraph()  # 여백

    # 메인 제목
    add_paragraph_styled(doc, "훈련운영계획서", 30, bold=True, color=DARK_GRAY,
                         alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=20)

    # 설명 문구
    desc_text = ("본 운영계획서는 심사지표의 평가 취지에 맞게 기관의 역량과 운영 체계가 어떻게 "
                 "구성·운영되는지를 이해하기 위한 자료입니다.\n"
                 "각 항목은 상호 연계된 하나의 운영 구조로 인식하고 작성해 주시기 바랍니다.\n"
                 "작성 분량은 최대 50면으로 제한합니다.")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    for line in desc_text.split('\n'):
        run = p.add_run(line)
        set_run_font(run, 11, color=LIGHT_GRAY)
        p.add_run('\n')

    # 하단 날짜 및 기관명
    for _ in range(4):
        doc.add_paragraph()

    add_paragraph_styled(doc, "{{작성일자}}", 21, color=MED_GRAY,
                         alignment=WD_ALIGN_PARAGRAPH.LEFT)
    add_paragraph_styled(doc, "{{기관명}}", 24, bold=True, color=MED_GRAY,
                         alignment=WD_ALIGN_PARAGRAPH.LEFT)

    add_page_break(doc)

    # ════════════════════════════════════════
    # 목차 (Table of Contents)
    # ════════════════════════════════════════
    add_paragraph_styled(doc, "목  차", 21, bold=True, color=DARK_GRAY,
                         alignment=WD_ALIGN_PARAGRAPH.LEFT, space_after=12)

    toc_items = [
        ("Ⅰ.", "사업 목적 및 인력 양성 계획", [
            ("1.", "사업 목표 및 추진 계획", ["사업 목표", "사업 추진 계획"]),
            ("2.", "AI 역량 기반의 인력 양성 계획", ["AI 전문 역량", "AI 인력 양성 계획"]),
        ]),
        ("Ⅱ.", "훈련인프라 확보 및 활용 계획", [
            ("1.", "조직 및 인적자원 운영 체계", [
                "조직 및 인적자원 확보", "AI 전문인력 활용 관리", "훈련운영인력 활용 관리"
            ]),
            ("2.", "시설·장비 및 AI 운영 체계", [
                "시설 및 장비 확보", "시설 및 장비 활용 관리", "AI 윤리 및 보안 거버넌스"
            ]),
        ]),
        ("Ⅲ.", "훈련 운영 관리 계획", [
            ("1.", "훈련 운영 및 학습 관리 체계", []),
            ("2.", "훈련 성과 관리 체계", []),
        ]),
        ("Ⅳ.", "훈련과정별 운영 개요 및 직업훈련 수료증 양식", []),
    ]

    for chapter_num, chapter_title, sections in toc_items:
        # 장
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = Pt(24)
        run_ch = p.add_run(f"{chapter_num} ")
        set_run_font(run_ch, 15, bold=True, color=DARK_GRAY)
        run_title = p.add_run(chapter_title)
        set_run_font(run_title, 15, bold=True, color=DARK_GRAY)

        for sec_item in sections:
            if isinstance(sec_item, tuple):
                sec_num, sec_title, subs = sec_item
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1.0)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = Pt(20)
                run = p.add_run(f"{sec_num} {sec_title}")
                set_run_font(run, 13, color=DARK_GRAY)

                for sub in subs:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(2.0)
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = Pt(18)
                    run = p.add_run(f"  {sub}")
                    set_run_font(run, 13, color=DARK_GRAY)

    add_page_break(doc)

    # ════════════════════════════════════════
    # 참고. 심사항목 참조표
    # ════════════════════════════════════════
    add_paragraph_styled(doc, "참고. 심사항목 참조표", 15, bold=True, color=WHITE,
                         alignment=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)

    p = doc.add_paragraph()
    run = p.add_run("훈련운영계획서 작성 시 절차별 심사 항목과 요소를 고려하여 작성")
    set_run_font(run, 13, color=DARK_GRAY)

    ref_headers = ["심사절차", "심사 항목", "항목이동"]
    ref_rows = [
        ["1차 (서면)", "운영가능성", ""],
        ["", "운영실적 관련성", "운영실적"],
        ["", "운영계획 부합성", "운영계획"],
        ["", "사업 운영을 위한 조직 구성", "조직구성"],
        ["", "AI 훈련시설", "훈련시설"],
        ["", "AI 훈련장비", "훈련장비"],
        ["", "AI 운영 체계 구축", "AI운영체계"],
        ["2차 (인터뷰)", "사업 타당성", "사업타당성"],
        ["", "실행 가능성", "실행가능성"],
        ["", "AI 전문성", "AI 전문성"],
        ["", "AI 전문인력", "AI 전문인력"],
        ["", "훈련운영인력", "운영인력"],
        ["", "시설 및 장비", "시설·장비"],
        ["", "훈련 계획", "훈련계획"],
        ["", "훈련 운영", "훈련운영"],
        ["", "훈련 성과 관리", "성과관리"],
        ["부록", "과정운영계획", "과정계획"],
    ]
    create_styled_table(doc, ref_headers, ref_rows)

    add_page_break(doc)

    # ════════════════════════════════════════
    # Ⅰ. 사업 목적 및 인력 양성 계획
    # ════════════════════════════════════════

    # --- 장 제목 (배경색 블록) ---
    def add_chapter_title(doc, text):
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        set_cell_shading(cell, HEADER_BG)
        set_cell_text(cell, text, size_pt=26, bold=True, color=WHITE,
                      alignment=WD_ALIGN_PARAGRAPH.LEFT)
        doc.add_paragraph()  # 여백

    add_chapter_title(doc, "Ⅰ. 사업 목적 및 인력 양성 계획")

    # --- 1. 사업 목표 및 추진 계획 ---
    def add_section_title(doc, text):
        """1. xxx 형태의 절 제목"""
        add_paragraph_styled(doc, text, 15, bold=True, color=DARK_GRAY,
                             space_before=12, space_after=6)

    def add_subsection_title(doc, text):
        """1.1 xxx 형태의 소절 제목"""
        add_paragraph_styled(doc, text, 13, bold=True, color=DARK_GRAY,
                             space_before=10, space_after=4)

    def add_related_question(doc, text):
        """관련 문항 박스"""
        t = doc.add_table(rows=1, cols=1)
        cell = t.rows[0].cells[0]
        set_cell_shading(cell, SUBHEADER_BG)
        set_cell_text(cell, text, size_pt=10, bold=False, color=ACCENT_BLUE,
                      alignment=WD_ALIGN_PARAGRAPH.LEFT)
        doc.add_paragraph()

    def add_checklist_title(doc, text):
        """□ 형태 체크리스트 제목"""
        add_paragraph_styled(doc, f"□ {text}", 12, bold=True, color=DARK_GRAY,
                             space_before=8, space_after=4)

    def add_sub_checklist(doc, text):
        """ 형태 하위 체크리스트"""
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(f"  {text}")
        set_run_font(run, 11, color=DARK_GRAY)

    def add_body_text(doc, text):
        """본문 텍스트"""
        add_paragraph_styled(doc, text, 11, color=DARK_GRAY,
                             space_before=2, space_after=2)

    def add_placeholder(doc, text):
        """플레이스홀더 텍스트"""
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f"{{{{{text}}}}}")
        set_run_font(run, 11, color=MED_GRAY)

    # ── 1. 사업 목표 및 추진 계획 ──
    add_section_title(doc, "1. 사업 목표 및 추진 계획")

    add_subsection_title(doc, "1.1 사업 목표")
    add_related_question(doc, "[2차] 사업 목적과 주요 내용이 정책 방향 및 사업 목표에 부합하는가?\n"
                              "  • 정부 AI 인재양성 정책 방향 및 AI 캠퍼스 사업 목표와의 정합성\n"
                              "  • 사업 추진 근거·필요성 및 사업 범위·대상·성과 목표 설정의 타당성\n"
                              "  • 훈련과정 편성 방향이 정책 방향 및 사업 목표에 부합하는지 여부")

    add_checklist_title(doc, "정부 AI 인재양성 정책을 반영한 사업 추진 기본 방향")
    add_sub_checklist(doc, "1) 정부 AI 인재양성 정책과의 정합성")
    add_placeholder(doc, "정부 비전 및 목표, 본 기관의 목표를 기재")

    add_sub_checklist(doc, "2) 본 기관의 전략적 대응 방향")
    add_placeholder(doc, "산업 수요 분석 및 전략적 대응 방안 기재")

    add_sub_checklist(doc, "3) 사업 범위 및 대상")
    add_placeholder(doc, "훈련 직군, 양성 규모, 훈련 기간 기재")

    # 훈련과정 표 템플릿
    add_paragraph_styled(doc, "< 훈련과정 개요 >", 10, bold=True, color=ACCENT_BLUE,
                         space_before=8)
    course_headers = ["직군", "훈련과정", "훈련시간", "연간 양성인원"]
    course_rows = [
        ["{{직군1}}", "{{과정명1}}", "{{시간1}}", "{{인원1}}"],
        ["{{직군2}}", "{{과정명2}}", "{{시간2}}", "{{인원2}}"],
        ["{{직군3}}", "{{과정명3}}", "{{시간3}}", "{{인원3}}"],
        ["합계", "", "", "{{총인원}}"],
    ]
    create_styled_table(doc, course_headers, course_rows)

    add_sub_checklist(doc, "4) 성과 목표")
    add_placeholder(doc, "취업률, 수료율, 모집률, 훈련생 만족도 등 성과 목표 기재")

    add_checklist_title(doc, "훈련과정 편성 방향의 정책 및 사업 부합성")
    add_sub_checklist(doc, "1) 정부 직군 체계와의 정합성")
    add_placeholder(doc, "직군별 훈련과정 매핑 내용 기재")
    add_sub_checklist(doc, "2) 산업 트렌드 및 기술 동향 반영")
    add_placeholder(doc, "산업 트렌드 반영 내용 기재")
    add_sub_checklist(doc, "3) 실무 역량 중심 교육 설계")
    add_placeholder(doc, "PBL, 프로젝트 생애주기 등 실무 교육 설계 기재")
    add_sub_checklist(doc, "4) 훈련과정별 주요 역량 목표")
    add_placeholder(doc, "과정별 주요 역량 목표 기재")

    # ── 1.2 사업 추진 계획 ──
    add_subsection_title(doc, "1.2 사업 추진 계획")
    add_related_question(doc, "[2차] 사업 목표 달성을 위한 추진계획이 유기적으로 설계되어 있는가?\n"
                              "  • 추진 일정·단계별 수행계획의 구체성과 현실성\n"
                              "  • 훈련 규모에 부합하는 인력 및 자원 배분의 적절성\n"
                              "  • 훈련기관-참여기업 간 역할 분담 및 협력 구조의 실효성")

    add_checklist_title(doc, "추진 일정 및 단계별 수행 계획")
    add_sub_checklist(doc, "사업 및 훈련과정 전체 추진 일정 (연 단위)")
    add_placeholder(doc, "연차별 운영 규모 및 추진 일정 기재")

    # 운영 일정표 템플릿
    schedule_headers = ["직군", "과정명", "1기", "2기", "3기", "4기", "양성인원"]
    schedule_rows = [
        ["{{직군}}", "{{과정명}}", "{{일정}}", "{{일정}}", "{{일정}}", "{{일정}}", "{{인원}}"],
    ]
    create_styled_table(doc, schedule_headers, schedule_rows)

    add_sub_checklist(doc, "단계별 주요 수행 내용 (기획-모집-운영-성과관리)")
    # 단계별 테이블
    phase_headers = ["단계", "기간", "핵심 활동", "핵심 산출물"]
    phase_rows = [
        ["기획", "D-60 ~ D-30", "{{수행내용}}", "{{산출물}}"],
        ["모집", "D-60 ~ D-0", "{{수행내용}}", "{{산출물}}"],
        ["운영", "D-0 ~ D+180", "{{수행내용}}", "{{산출물}}"],
        ["성과관리", "D+180 ~ D+360", "{{수행내용}}", "{{산출물}}"],
    ]
    create_styled_table(doc, phase_headers, phase_rows)

    # 각 단계 상세 템플릿
    for phase in ["1단계: 기획", "2단계: 모집", "3단계: 운영", "4단계: 성과관리"]:
        add_paragraph_styled(doc, f"▶ {phase}", 10.6, bold=True, color=ACCENT_BLUE,
                             space_before=8)
        detail_headers = ["구분", "수행 내용", "세부 활동", "담당", "산출물"]
        detail_rows = [
            ["{{구분}}", "{{수행내용}}", "{{세부활동}}", "{{담당}}", "{{산출물}}"],
        ]
        create_styled_table(doc, detail_headers, detail_rows)

    add_checklist_title(doc, "추진 일정에 따른 인력 및 자원 분배 계획")
    add_sub_checklist(doc, "훈련 규모 대비 인력 투입 계획")
    resource_headers = ["구분", "역할", "1년차", "2년차", "3년차", "비고"]
    resource_rows = [
        ["AI전문인력", "{{역할}}", "{{인원}}", "{{인원}}", "{{인원}}", "{{비고}}"],
        ["훈련운영인력", "{{역할}}", "{{인원}}", "{{인원}}", "{{인원}}", "{{비고}}"],
    ]
    create_styled_table(doc, resource_headers, resource_rows)

    add_sub_checklist(doc, "단계별 인력 및 자원 배분 방식")
    add_sub_checklist(doc, "운영 중 자원 조정 및 보완 계획")

    add_checklist_title(doc, "훈련기관-참여기업 간 역할 분담 및 협력 구조")
    add_sub_checklist(doc, "참여기업 목록 및 훈련 참여 방식")
    company_headers = ["기업명", "주요 보유 기술 및 서비스", "주요 참여 방식 (역할)"]
    company_rows = [
        ["{{기업명}}", "{{보유기술}}", "{{참여방식}}"],
        ["{{기업명}}", "{{보유기술}}", "{{참여방식}}"],
    ]
    create_styled_table(doc, company_headers, company_rows)

    add_sub_checklist(doc, "훈련기관과 참여기업의 역할 분담")
    role_headers = ["구분", "훈련기관", "참여기업"]
    role_rows = [
        ["기획 및 모집", "{{훈련기관 역할}}", "{{참여기업 역할}}"],
        ["훈련 운영", "{{훈련기관 역할}}", "{{참여기업 역할}}"],
        ["인프라 지원", "{{훈련기관 역할}}", "{{참여기업 역할}}"],
        ["성과 관리", "{{훈련기관 역할}}", "{{참여기업 역할}}"],
    ]
    create_styled_table(doc, role_headers, role_rows)

    add_sub_checklist(doc, "협력 과정에서의 의사소통 관리 방식")
    comm_headers = ["소통 방식", "세부 내용", "주기"]
    comm_rows = [
        ["단계별 데이터 인계", "{{내용}}", "{{주기}}"],
        ["실시간 일정 및 이슈 관리", "{{내용}}", "{{주기}}"],
        ["정기 성과 점검 및 환류", "{{내용}}", "{{주기}}"],
    ]
    create_styled_table(doc, comm_headers, comm_rows)

    add_page_break(doc)

    # ── 2. AI 역량 기반의 인력 양성 계획 ──
    add_section_title(doc, "2. AI 역량 기반의 인력 양성 계획")

    add_subsection_title(doc, "2.1 AI 전문 역량")
    add_related_question(doc, "[1차] 훈련기관의 최근 3년간 AI 분야 운영 실적이 AI 캠퍼스 사업 수행 역량을 입증하는가?\n"
                              "[2차] 훈련기관 및 참여기업은 AI 분야 전문성을 보유하고 있는가?")

    add_checklist_title(doc, "훈련기관의 최근 3년간 AI 분야 훈련과정 운영 실적")
    add_sub_checklist(doc, "AI 분야 고용노동부 지원 사업 훈련과정 운영 실적 (KDT/KDT 외)")
    track_record_headers = ["사업명", "훈련과정명", "직무", "대상자", "수준", "시간", "누적 인원", "주요성과"]
    track_record_rows = [
        ["{{사업명}}", "{{과정명}}", "{{직무}}", "{{대상}}", "{{수준}}",
         "{{시간}}", "{{인원}}", "{{성과}}"],
    ]
    create_styled_table(doc, track_record_headers, track_record_rows)

    add_sub_checklist(doc, "AI 분야 고용노동부 지원 사업 외 훈련과정 운영 실적")
    add_sub_checklist(doc, "훈련기관의 AI 분야 연구·과업·사업 실적")
    add_sub_checklist(doc, "참여기업의 AI 분야 교육·연구·실무협력 성과")

    add_subsection_title(doc, "2.2 AI 인력 양성 계획")
    add_sub_checklist(doc, "수요 분석 결과를 반영한 인력양성목표 및 훈련과정 로드맵")
    add_placeholder(doc, "산업수요 및 직무분석, 인력양성목표 기재")

    add_sub_checklist(doc, "훈련과정별 직무 수행 모델 및 역량 체계")
    add_sub_checklist(doc, "훈련과정별 상세 교육과정 (커리큘럼)")

    add_page_break(doc)

    # ════════════════════════════════════════
    # Ⅱ. 훈련인프라 확보 및 활용 계획
    # ════════════════════════════════════════
    add_chapter_title(doc, "Ⅱ. 훈련인프라 확보 및 활용 계획")

    add_section_title(doc, "1. 조직 및 인적자원 운영 체계")

    add_subsection_title(doc, "1.1 조직 및 인적자원 확보")
    add_related_question(doc, "[1차] 사업 운영을 위한 조직이 적절히 구성되어 있는가?")
    add_checklist_title(doc, "훈련 운영 조직 체계")
    add_placeholder(doc, "조직 구성도 및 부서별 역할 기재")

    add_subsection_title(doc, "1.2 AI 전문인력 활용 관리")
    add_related_question(doc, "[2차] AI 전문인력이 적절히 확보 및 활용되는가?")
    add_placeholder(doc, "AI 전문인력 현황 및 활용 계획 기재")

    add_subsection_title(doc, "1.3 훈련운영인력 활용 관리")
    add_placeholder(doc, "훈련운영인력 현황 및 활용 계획 기재")

    add_section_title(doc, "2. 시설·장비 및 AI 운영 체계")

    add_subsection_title(doc, "2.1 시설 및 장비 확보")
    add_related_question(doc, "[1차] AI 훈련시설 및 장비가 적절히 확보되어 있는가?")

    add_checklist_title(doc, "훈련시설 확보 현황")
    facility_headers = ["구분", "시설명", "규모", "용도", "확보 방식"]
    facility_rows = [
        ["{{구분}}", "{{시설명}}", "{{규모}}", "{{용도}}", "{{확보방식}}"],
    ]
    create_styled_table(doc, facility_headers, facility_rows)

    add_checklist_title(doc, "훈련장비 확보 현황")
    equip_headers = ["구분", "장비명", "규격/사양", "수량", "용도", "확보 방식"]
    equip_rows = [
        ["{{구분}}", "{{장비명}}", "{{규격}}", "{{수량}}", "{{용도}}", "{{확보방식}}"],
    ]
    create_styled_table(doc, equip_headers, equip_rows)

    add_subsection_title(doc, "2.2 시설 및 장비 활용 관리")
    add_placeholder(doc, "시설 및 장비 활용 관리 계획 기재")

    add_subsection_title(doc, "2.3 AI 윤리 및 보안 거버넌스")
    add_placeholder(doc, "AI 윤리 및 보안 거버넌스 구축 계획 기재")

    add_page_break(doc)

    # ════════════════════════════════════════
    # Ⅲ. 훈련 운영 관리 계획
    # ════════════════════════════════════════
    add_chapter_title(doc, "Ⅲ. 훈련 운영 관리 계획")

    add_section_title(doc, "1. 훈련 운영 및 학습 관리 체계")
    add_related_question(doc, "[2차] 훈련 계획 및 훈련 운영이 체계적으로 수립되어 있는가?")

    add_checklist_title(doc, "훈련 운영 체계")
    add_placeholder(doc, "훈련 운영 및 학습 관리 체계 전반 기재")

    add_checklist_title(doc, "학습 관리 체계")
    add_placeholder(doc, "LMS, 학습 진도 관리, 훈련생 케어 체계 기재")

    add_checklist_title(doc, "교강사 운영 체계")
    add_placeholder(doc, "교강사 현황 및 운영 방식 기재")

    add_section_title(doc, "2. 훈련 성과 관리 체계")
    add_related_question(doc, "[2차] 훈련 성과 관리 체계가 체계적으로 구축되어 있는가?")

    add_checklist_title(doc, "성과 관리 지표")
    perf_headers = ["성과지표명", "달성목표(1차)", "달성목표(2차)", "달성목표(3차)", "최종목표"]
    perf_rows = [
        ["{{지표명}}", "{{목표}}", "{{목표}}", "{{목표}}", "{{목표}}"],
    ]
    create_styled_table(doc, perf_headers, perf_rows)

    add_checklist_title(doc, "과정별 운영성과 분석 및 개선과제 도출 체계")
    add_placeholder(doc, "정량/정성 분석 체계, NPS, 환류 체계 기재")

    add_page_break(doc)

    # ════════════════════════════════════════
    # Ⅳ. 훈련과정별 운영 개요
    # ════════════════════════════════════════
    add_chapter_title(doc, "Ⅳ. 훈련과정별 운영 개요 및 직업훈련 수료증 양식")

    add_checklist_title(doc, "훈련과정별 운영 개요")
    add_placeholder(doc, "과정별 운영 개요 기재 (과정 수만큼 복제하여 작성)")

    # 과정별 템플릿
    for proc_num in range(1, 4):
        add_paragraph_styled(doc, f"▶ 훈련과정 {proc_num}", 12, bold=True,
                             color=ACCENT_BLUE, space_before=12)

        proc_headers = ["훈련과정명", "AI 직군", "훈련 교‧강사"]
        proc_rows = [
            ["{{과정명}}", "{{직군}}", ""],
        ]
        create_styled_table(doc, proc_headers, proc_rows)

        # 교강사 상세 테이블
        instructor_headers = ["연번", "구분", "역할", "주요업무", "투입시간", "성명"]
        instructor_rows = [
            ["1", "내부/외부", "{{역할}}", "{{주요업무}}", "{{시간}}", "{{성명}}"],
            ["2", "내부/외부", "{{역할}}", "{{주요업무}}", "{{시간}}", "{{성명}}"],
        ]
        create_styled_table(doc, instructor_headers, instructor_rows)

    # 직업훈련 수료증 양식
    add_page_break(doc)
    add_checklist_title(doc, "직업훈련 수료증 양식")
    add_body_text(doc, "※ 필수기재사항: 훈련생 성명, 생년월일, 훈련과정명, 훈련기간, "
                       "획득한 직무 역량(Skill-set), 주관기관 연락처")

    cert_headers = ["항목", "내용"]
    cert_rows = [
        ["훈련생 성명", "{{성명}}"],
        ["생년월일", "{{생년월일}}"],
        ["훈련과정명", "{{과정명}}"],
        ["훈련기간", "{{기간}}"],
        ["획득 직무 역량", "{{역량}}"],
        ["주관기관", "{{기관명}}"],
        ["연락처", "{{연락처}}"],
    ]
    create_styled_table(doc, cert_headers, cert_rows)

    # 부록: 훈련과정 개요서
    add_page_break(doc)
    add_paragraph_styled(doc, "참고. 훈련과정 개요서(예시)", 18, bold=True,
                         color=DARK_GRAY, space_after=8)

    overview_items = [
        ("과정명", "{{과정명}}"),
        ("주관기관", "{{기관명}}"),
        ("훈련기간", "총 {{개월}}개월"),
        ("훈련시간", "총 {{시간}}시간 (집체 {{시간}}시간, 실시간 비대면 {{시간}}시간)"),
        ("훈련목표", "{{훈련목표}}"),
    ]
    overview_headers = ["항목", "내용"]
    create_styled_table(doc, overview_headers, overview_items)

    add_paragraph_styled(doc, "훈련이 목표로 하는 핵심 직무 역량", 12, bold=True,
                         color=DARK_GRAY, space_before=8)
    add_placeholder(doc, "핵심 직무 역량 내용 기재")

    add_paragraph_styled(doc, "교과과정 개요", 12, bold=True,
                         color=DARK_GRAY, space_before=8)
    curriculum_headers = ["모듈", "교과목", "시간", "주요 학습내용", "담당강사"]
    curriculum_rows = [
        ["{{모듈명}}", "{{교과목}}", "{{시간}}", "{{학습내용}}", "{{강사}}"],
    ]
    create_styled_table(doc, curriculum_headers, curriculum_rows)

    # ── 저장 ──
    output_dir = os.path.dirname(os.path.abspath(__file__))
    output_path = os.path.join(output_dir, "aicampus_proposal_template.docx")
    doc.save(output_path)
    print(f"템플릿 생성 완료: {output_path}")
    return output_path


if __name__ == "__main__":
    path = generate_template()
    print(f"Saved: {path}")
