# -*- coding: utf-8 -*-
"""
도쿄일렉트론코리아 RAG 심화 교육 상세 일정 Excel 생성기
3일 × 8H = 24H, 09:00-18:00 (점심 1H 포함)
"""
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if sys.platform.startswith("win"):
    os.environ.setdefault("PYTHONIOENCODING", "utf-8")

# ===== Style helpers =====
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
DAY_FILL = PatternFill("solid", fgColor="DEEBF7")
DAY_FONT = Font(name="맑은 고딕", size=11, bold=True, color="1F4E78")
BREAK_FILL = PatternFill("solid", fgColor="F2F2F2")
BREAK_FONT = Font(name="맑은 고딕", size=10, italic=True, color="7F7F7F")
LUNCH_FILL = PatternFill("solid", fgColor="FFF2CC")
LUNCH_FONT = Font(name="맑은 고딕", size=10, italic=True, color="806000")
BODY_FONT = Font(name="맑은 고딕", size=10)
WRAP_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


# ===== Curriculum data =====
DAYS = [
    {
        "label": "1일차",
        "date": "2026-05-13 (수)",
        "title": "RAG 기초 및 데이터 처리",
        "objective": "LangChain 기초·RAG 개념 이해부터 ChromaDB 기반 기본 RAG 파이프라인 구축, 청킹·Query Rewriting까지 완성",
        "blocks": [
            ("09:00-10:30", 90, "오리엔테이션 + LangChain 기초",
             "과정 안내 / 강사·수강생 소개\nLangChain 기본 구조 / 간단한 Chain 실습", "강의+실습"),
            ("10:30-10:45", 15, "휴식", "", "BREAK"),
            ("10:45-12:00", 75, "프롬프트 템플릿 활용",
             "PromptTemplate, ChatPromptTemplate, Few-shot 기본 사용법\nLCEL 기반 Q&A 챗봇 실습", "강의+실습"),
            ("12:00-13:00", 60, "점심", "", "LUNCH"),
            ("13:00-14:30", 90, "RAG 개념과 벡터 검색 원리",
             "RAG 개념 이해 — LLM 단독 사용의 한계, Retrieval + Generation 구조\n임베딩과 벡터 검색 원리", "강의"),
            ("14:30-14:45", 15, "휴식", "", "BREAK"),
            ("14:45-16:15", 90, "ChromaDB 벡터 저장소 구축",
             "ChromaDB 기반 벡터 저장소 구축·검색 실습\nChromaDB 기본 인덱싱 방식 HNSW 포함, 벡터 인덱싱 개념", "강의+실습"),
            ("16:15-16:30", 15, "휴식", "", "BREAK"),
            ("16:30-18:00", 90, "기본 RAG 파이프라인 + 청킹·Query Rewriting",
             "기본형 RAG 파이프라인 구성 실습 / Naive RAG 한계 진단\n고정 길이 / 재귀적 / 구조 기반 / 시맨틱 청킹 비교\nQuery Rewriting을 통한 검색 질문 개선 실습", "실습"),
        ],
    },
    {
        "label": "2일차",
        "date": "2026-05-14 (목)",
        "title": "검색 고도화 및 Advanced RAG 기법",
        "objective": "Query 확장·Hybrid Retrieval·Re-ranking·Context 압축·Multi-Representation 인덱싱까지 검색 품질 개선 기법 종합 실습",
        "blocks": [
            ("09:00-10:30", 90, "Query 확장 기법",
             "Multi-Query Retrieval — 하나의 질문을 여러 검색 쿼리로 확장\nHyDE — 가상 답변 생성을 통한 검색 품질 보완", "강의+실습"),
            ("10:30-10:45", 15, "휴식", "", "BREAK"),
            ("10:45-12:00", 75, "Semantic Routing + Hybrid Retrieval",
             "Semantic Routing — 질문 유형에 따른 검색 경로 선택\nHybrid Retrieval — BM25 Sparse Retrieval + Dense Retrieval 결합", "강의+실습"),
            ("12:00-13:00", 60, "점심", "", "LUNCH"),
            ("13:00-14:30", 90, "검색 결과 결합·재정렬",
             "RRF (Reciprocal Rank Fusion) — Sparse/Dense 검색 결과 순위 기반 결합\nCross-Encoder Re-ranking — 검색 후보 재정렬 실습", "강의+실습"),
            ("14:30-14:45", 15, "휴식", "", "BREAK"),
            ("14:45-16:15", 90, "Context Filtering & Compression",
             "검색 결과가 방대할 때 LLM 기반 요약·중복 필터링\n질문 기준 문맥 압축 실습", "강의+실습"),
            ("16:15-16:30", 15, "휴식", "", "BREAK"),
            ("16:30-18:00", 90, "Multi-Representation Indexing + 종합",
             "원문 외 요약/대표 표현을 함께 활용하는 인덱싱 패턴\nDay 2 종합 실습 및 정리", "실습"),
        ],
    },
    {
        "label": "3일차",
        "date": "2026-05-15 (금)",
        "title": "평가, 운영 모니터링, 지식그래프 기초",
        "objective": "평가셋 설계·RAGAS·Langfuse 기반 정량 평가 체계화, Vector RAG 한계 점검 후 지식그래프·미니 경진대회로 마무리",
        "blocks": [
            ("09:00-10:30", 90, "평가셋 설계 + Retrieval 평가",
             "평가셋 설계 — 질문, 기준 답변, 근거 문서 구성\nRetrieval 평가 — Hit@K, Recall@K 계산 실습", "강의+실습"),
            ("10:30-10:45", 15, "휴식", "", "BREAK"),
            ("10:45-12:00", 75, "Generation 평가 + RAGAS 지표 관계",
             "Generation 평가 — Faithfulness 중심의 LLM-as-a-Judge 평가\n수동 평가 지표와 RAGAS 지표의 관계", "강의+실습"),
            ("12:00-13:00", 60, "점심", "", "LUNCH"),
            ("13:00-14:30", 90, "RAGAS · Langfuse · TruLens",
             "RAGAS 입력 구조 — user_input, retrieved_contexts, response, reference\nLangfuse 기반 RAG 요청 Trace 확인\nRAGAS, Langfuse, TruLens의 역할 구분", "강의+실습"),
            ("14:30-14:45", 15, "휴식", "", "BREAK"),
            ("14:45-16:15", 90, "Vector RAG 한계 + 지식그래프 기초",
             "관계형 질문에서 Vector RAG가 흔들리는 지점 확인\n온톨로지 / 지식그래프 기초 — Entity, Relation, Triple 개념\n짧은 문장에서 트리플 추출 미니 실습\nGraph RAG로 확장하기 전 설계 체크리스트", "강의+실습"),
            ("16:15-16:30", 15, "휴식", "", "BREAK"),
            ("16:30-18:00", 90, "미니 경진대회 (최종 과제)",
             "강사 제공 ~190p 문서 기반 RAG 개선 과제\nLangfuse 기반 자동채점 평가 / 결과 발표 / 강사 피드백", "실습+평가"),
        ],
    },
]


# ===== Workbook setup =====
wb = Workbook()
wb.remove(wb.active)


# ===== Sheet 1: 교육 개요 =====
ws_overview = wb.create_sheet("교육 개요")
ws_overview.sheet_view.showGridLines = False

# Title
ws_overview["A1"] = "도쿄일렉트론코리아 RAG 심화 교육 — 상세 일정"
ws_overview["A1"].font = Font(name="맑은 고딕", size=16, bold=True, color="1F4E78")
ws_overview.merge_cells("A1:E1")
ws_overview.row_dimensions[1].height = 28

# Meta
meta_rows = [
    ("교육 과정", "RAG (Retrieval-Augmented Generation) 심화"),
    ("교육 일정", "2026년 5월 13일(수) ~ 5월 15일(금) · 3일간"),
    ("교육 시간", "09:00 - 18:00 (점심 1H 포함, 일 8H × 3일 = 총 24H)"),
    ("교육 장소", "도쿄일렉트론코리아 (협의)"),
    ("강사", "김민수"),
    ("교육 운영", "모두의연구소 (담당: 김영광 / yk.kim@modulabs.co.kr)"),
    ("실습 자료", "강사 제공 약 190페이지 분량 PDF 문서 (도메인 비종속). 미니 경진대회도 동일 자료 기반"),
    ("교육 난이도", "중급 — 개발자/일반 업무 담당자 혼합 수강 고려, 심화 주제는 참고 자료 제공"),
]
for i, (k, v) in enumerate(meta_rows, start=3):
    ws_overview.cell(row=i, column=1, value=k).font = Font(name="맑은 고딕", size=10, bold=True)
    ws_overview.cell(row=i, column=1).fill = PatternFill("solid", fgColor="F2F2F2")
    ws_overview.cell(row=i, column=1).alignment = WRAP_LEFT
    ws_overview.cell(row=i, column=2, value=v).font = BODY_FONT
    ws_overview.cell(row=i, column=2).alignment = WRAP_LEFT
    ws_overview.merge_cells(start_row=i, start_column=2, end_row=i, end_column=5)
    for col in range(1, 6):
        ws_overview.cell(row=i, column=col).border = BORDER

# Day-by-day summary
sum_start = 3 + len(meta_rows) + 2
ws_overview.cell(row=sum_start, column=1, value="일자별 학습 목표").font = Font(name="맑은 고딕", size=13, bold=True, color="1F4E78")
ws_overview.merge_cells(start_row=sum_start, start_column=1, end_row=sum_start, end_column=5)

headers = ["차수", "일자", "주제", "학습 목표", "비고"]
for col, h in enumerate(headers, start=1):
    cell = ws_overview.cell(row=sum_start + 1, column=col, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = WRAP_CENTER
    cell.border = BORDER

for i, day in enumerate(DAYS):
    r = sum_start + 2 + i
    ws_overview.cell(row=r, column=1, value=day["label"]).alignment = WRAP_CENTER
    ws_overview.cell(row=r, column=2, value=day["date"]).alignment = WRAP_CENTER
    ws_overview.cell(row=r, column=3, value=day["title"]).alignment = WRAP_LEFT
    ws_overview.cell(row=r, column=4, value=day["objective"]).alignment = WRAP_LEFT
    ws_overview.cell(row=r, column=5, value="8H").alignment = WRAP_CENTER
    for col in range(1, 6):
        ws_overview.cell(row=r, column=col).font = BODY_FONT
        ws_overview.cell(row=r, column=col).border = BORDER
    ws_overview.row_dimensions[r].height = 60

# Column widths for overview
ws_overview.column_dimensions["A"].width = 14
ws_overview.column_dimensions["B"].width = 22
ws_overview.column_dimensions["C"].width = 32
ws_overview.column_dimensions["D"].width = 60
ws_overview.column_dimensions["E"].width = 10


# ===== Sheet 2: 상세 일정 (통합) =====
ws_detail = wb.create_sheet("상세 일정")
ws_detail.sheet_view.showGridLines = False

ws_detail["A1"] = "전체 상세 일정 (1·2·3일차)"
ws_detail["A1"].font = Font(name="맑은 고딕", size=14, bold=True, color="1F4E78")
ws_detail.merge_cells("A1:F1")
ws_detail.row_dimensions[1].height = 24

headers = ["차수", "시간", "시간(분)", "모듈", "세부 학습 내용", "진행 형식"]
for col, h in enumerate(headers, start=1):
    cell = ws_detail.cell(row=2, column=col, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = WRAP_CENTER
    cell.border = BORDER
ws_detail.row_dimensions[2].height = 22

row = 3
for day in DAYS:
    # Day separator row
    label = f"{day['label']} | {day['date']} | {day['title']} (8H)"
    cell = ws_detail.cell(row=row, column=1, value=label)
    cell.font = DAY_FONT
    cell.fill = DAY_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    cell.border = BORDER
    ws_detail.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    for col in range(1, 7):
        ws_detail.cell(row=row, column=col).border = BORDER
    ws_detail.row_dimensions[row].height = 22
    row += 1

    for time_range, minutes, module, detail, fmt in day["blocks"]:
        ws_detail.cell(row=row, column=1, value=day["label"])
        ws_detail.cell(row=row, column=2, value=time_range)
        ws_detail.cell(row=row, column=3, value=minutes)
        ws_detail.cell(row=row, column=4, value=module)
        ws_detail.cell(row=row, column=5, value=detail)
        ws_detail.cell(row=row, column=6, value=fmt)

        is_break = fmt == "BREAK"
        is_lunch = fmt == "LUNCH"
        for col in range(1, 7):
            c = ws_detail.cell(row=row, column=col)
            c.border = BORDER
            if is_break:
                c.fill = BREAK_FILL
                c.font = BREAK_FONT
            elif is_lunch:
                c.fill = LUNCH_FILL
                c.font = LUNCH_FONT
            else:
                c.font = BODY_FONT
            if col in (1, 2, 3, 6):
                c.alignment = WRAP_CENTER
            else:
                c.alignment = WRAP_LEFT

        # If break/lunch, blank module/detail/format presentation
        if is_break:
            ws_detail.cell(row=row, column=4, value="휴식")
            ws_detail.cell(row=row, column=5, value="")
            ws_detail.cell(row=row, column=6, value="-")
        elif is_lunch:
            ws_detail.cell(row=row, column=4, value="점심")
            ws_detail.cell(row=row, column=5, value="")
            ws_detail.cell(row=row, column=6, value="-")

        # Row height: detail rows higher
        if is_break or is_lunch:
            ws_detail.row_dimensions[row].height = 18
        else:
            line_count = max(detail.count("\n") + 1, 2)
            ws_detail.row_dimensions[row].height = max(40, line_count * 18)
        row += 1

# Column widths for detail
ws_detail.column_dimensions["A"].width = 8
ws_detail.column_dimensions["B"].width = 14
ws_detail.column_dimensions["C"].width = 9
ws_detail.column_dimensions["D"].width = 30
ws_detail.column_dimensions["E"].width = 70
ws_detail.column_dimensions["F"].width = 12

# Freeze top header
ws_detail.freeze_panes = "A3"


# ===== Save =====
out_path = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "도쿄일렉트론코리아_RAG교육_상세일정.xlsx"
)
wb.save(out_path)
print(f"Saved: {out_path}")
