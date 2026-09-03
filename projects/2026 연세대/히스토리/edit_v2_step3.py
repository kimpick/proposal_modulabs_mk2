"""
Step 3: 마스터 시트 기반 추가 보완
- RFP 공식 3대 목표 + 우리 응답 매핑
- 캠퍼스·학기별 운영 규모 상세 표 추가
- 교육 커리큘럼 핵심 주제 영역 표 추가
"""
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from copy import deepcopy

PATH = '연세AX_제안서_모두의연구소_v2_260506.docx'
doc = Document(PATH)


def find_para_by_text(doc, contains_text, start_idx=0):
    for i, p in enumerate(doc.paragraphs):
        if i < start_idx:
            continue
        if contains_text in p.text:
            return i
    return -1


def insert_paragraph_after(reference_element, text, bold=False):
    new_p = OxmlElement('w:p')
    reference_element.addnext(new_p)
    r = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    if bold:
        b = OxmlElement('w:b')
        rpr.append(b)
        b2 = OxmlElement('w:bCs')
        rpr.append(b2)
    rfonts = OxmlElement('w:rFonts')
    rfonts.set(qn('w:eastAsia'), '맑은 고딕')
    rpr.append(rfonts)
    r.append(rpr)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    new_p.append(r)
    return new_p


def insert_table_after(reference_element, data, bold_header=True):
    """단순 표를 reference_element 뒤에 삽입. data는 list of list of str."""
    rows_n = len(data)
    cols_n = max(len(row) for row in data) if data else 0

    tbl = OxmlElement('w:tbl')

    # tblPr (기본 속성)
    tblPr = OxmlElement('w:tblPr')
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '5000')
    tblW.set(qn('w:type'), 'pct')
    tblPr.append(tblW)
    # 테두리
    tblBorders = OxmlElement('w:tblBorders')
    for border_type in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement(f'w:{border_type}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), '000000')
        tblBorders.append(b)
    tblPr.append(tblBorders)
    tbl.append(tblPr)

    # tblGrid
    tblGrid = OxmlElement('w:tblGrid')
    for _ in range(cols_n):
        gridCol = OxmlElement('w:gridCol')
        tblGrid.append(gridCol)
    tbl.append(tblGrid)

    # rows
    for r_idx, row_data in enumerate(data):
        tr = OxmlElement('w:tr')
        for c_idx in range(cols_n):
            tc = OxmlElement('w:tc')
            tcPr = OxmlElement('w:tcPr')
            tc.append(tcPr)
            p = OxmlElement('w:p')
            pPr = OxmlElement('w:pPr')
            p.append(pPr)
            r = OxmlElement('w:r')
            rPr = OxmlElement('w:rPr')
            if r_idx == 0 and bold_header:
                b = OxmlElement('w:b')
                rPr.append(b)
                b2 = OxmlElement('w:bCs')
                rPr.append(b2)
            rfonts = OxmlElement('w:rFonts')
            rfonts.set(qn('w:eastAsia'), '맑은 고딕')
            rPr.append(rfonts)
            r.append(rPr)
            t = OxmlElement('w:t')
            t.text = row_data[c_idx] if c_idx < len(row_data) else ''
            t.set(qn('xml:space'), 'preserve')
            r.append(t)
            p.append(r)
            tc.append(p)
            tr.append(tc)
        tbl.append(tr)

    reference_element.addnext(tbl)
    return tbl


# ====================================================================
# 1. RFP 공식 3대 목표 + 우리 응답 매핑 (사업 수행 목표 위에)
# ====================================================================
# "사업 수행 목표 및 추진 전략" 헤딩 뒤, "본 사업의 수행 목표를..." 단락 위에 삽입
heading_idx = -1
for i, p in enumerate(doc.paragraphs):
    if "사업 수행 목표" in p.text and "추진 전략" in p.text:
        heading_idx = i
        break

if heading_idx >= 0:
    anchor = doc.paragraphs[heading_idx]._element
    # 역순으로 삽입 (anchor 바로 다음에 차곡차곡)
    paras_to_insert = [
        ("발주처가 제안요청서에서 제시한 본 사업의 3대 목표는 다음과 같습니다.", False),
        ("① AI 전환 시대 대응 통합 교육 체계 구축 — 미래형 융합 인재 양성", True),
        ("② 핵심 AX 기술 기반 문제 해결 중심 실전형 교육 경험 제공", True),
        ("③ 수요 기반 맞춤형 교육과 지속적 품질 관리를 통한 교육 효과 극대화", True),
        ("", False),
        ("모두의연구소는 위 3대 목표에 응답하기 위해, 다음 네 가지 수행 목표를 설정하고 정량적 달성 기준을 제시합니다. 가·나는 ①·②에, 다는 ②·③에, 라는 ②에 직접 대응합니다.", False),
        ("", False),
    ]
    # reverse 순서로 삽입해야 결과적으로 순방향 됨
    for text, bold in reversed(paras_to_insert):
        insert_paragraph_after(anchor, text, bold=bold)
    print(f"INSERTED RFP 3대 목표 매핑 ({len(paras_to_insert)} 단락)")

# 기존 "본 사업의 수행 목표를 다음 네 가지로..." 단락 텍스트 살짝 수정 (중복 제거)
idx = find_para_by_text(doc, "본 사업의 수행 목표를 다음 네 가지로 설정하고")
if idx >= 0:
    # 그대로 유지 (전체 흐름상 자연스러움)
    pass

# ====================================================================
# 2. 캠퍼스·학기별 운영 규모 상세 표 추가
#    위치: "▶ 교육 커리큘럼 구성 방안 (총 90건...)" 표 다음
# ====================================================================
# 해당 표(curriculum table)를 찾기
curriculum_table = None
for t in doc.tables:
    if len(t.rows) > 0:
        first = [c.text.strip() for c in t.rows[0].cells]
        if '구분' in first and '내용' in first and '건수' in first and '주요 내용' in first:
            curriculum_table = t
            break

if curriculum_table is not None:
    anchor = curriculum_table._element
    # 표 위에 헤더 1줄 + 빈 줄 + 표 + 빈 줄 형태로 삽입
    # 순서: 헤더 → 표 (역순으로 anchor.addnext)

    # 캠퍼스·학기별 상세 표 데이터 (TSV 마스터시트 line 65-95 기반)
    detail_data = [
        ['캠퍼스', '학기', '구분', '과정명', '교육 횟수(회)', '시수(시간)', '총 시수(H)'],
        ['신촌캠퍼스 도서관', '2026-1학기', '체험 교육', '실감미디어 몰입 체험 정규 프로그램', '102', '1', '102'],
        ['', '', '세미나', 'AX 트렌드 세미나(상반기)', '2', '2', '4'],
        ['', '', '워크숍', 'AX 기술 실습 워크숍', '5', '4', '20'],
        ['', '', '멘토링', '프로젝트 기반 심화 교육 (2026 여름 AX 캠프)', '20', '4', '80'],
        ['', '', '소계', '', '129', '-', '206'],
        ['', '2026-2학기', '체험 교육', '실감미디어 몰입 체험 정규 프로그램', '150', '1', '150'],
        ['', '', '', '실감미디어 몰입 체험 특별 프로그램 I·II', '2', '-', '-'],
        ['', '', '세미나', 'AX 트렌드 세미나(하반기)', '2', '2', '4'],
        ['', '', '워크숍', 'AX 기술 실습 워크숍', '5', '4', '20'],
        ['', '', '멘토링', 'AI 활용 문제 해결 성과 발표 및 공유 대회', '2', '3', '6'],
        ['', '', '소계', '', '161', '-', '180'],
        ['국제캠퍼스 도서관', '2026-1학기', '워크숍', '생성형 AI 원데이 클래스', '10', '2', '20'],
        ['', '', '', '미디어 제작·활용 워크숍 (CapCut, PPT 등)', '5', '3', '15'],
        ['', '', '', 'AX 북클럽 공모전 연계 미디어 워크숍', '3', '3', '9'],
        ['', '', '멘토링', '공모전을 위한 기술 워크숍', '2', '8', '16'],
        ['', '', '소계', '', '20', '-', '60'],
        ['', '2026-2학기', '세미나', 'AI 전환 트렌드 세미나', '3', '2', '6'],
        ['', '', '워크숍', '생성형 AI 원데이 클래스', '5', '3', '15'],
        ['', '', '', '미디어 제작·활용 워크숍', '5', '3', '15'],
        ['', '', '', '공모전 연계 코딩 및 생성형 AI 워크숍', '10', '3', '30'],
        ['', '', '', 'AX 북클럽 공모전 연계 미디어 워크숍', '3', '3', '9'],
        ['', '', '멘토링', '공모전을 위한 기술 워크숍', '2', '8', '16'],
        ['', '', '소계', '', '28', '-', '91'],
        ['신촌캠퍼스 박물관', '연간', '인문학 세미나', 'AX 시대, 인문학의 미래 / 한글을 지킨 연세의 거인들', '3', '2', '6'],
        ['', '', 'AI 활용 워크숍', 'AI 창작 워크숍 (기초·심화)', '2', '3', '6'],
        ['', '', '소계', '', '5', '-', '12'],
        ['전체 합계', '', '', '', '343', '-', '549'],
    ]

    # 빈 단락 삽입 (표 다음 공백)
    insert_paragraph_after(anchor, "", bold=False)
    # 표 삽입 (anchor 바로 다음, 위에서 만든 빈 단락 앞에)
    insert_table_after(anchor, detail_data)
    # 표 위에 헤더 삽입 (anchor 바로 다음. 표보다 먼저 들어가야 함)
    # → 순서: 헤더 paragraph 먼저 삽입
    insert_paragraph_after(anchor, "▶ 캠퍼스·학기별 운영 규모 상세 (총 343회 549시간 산출 근거)", bold=True)
    print("INSERTED 캠퍼스·학기별 운영 규모 상세 표")

# ====================================================================
# 3. 교육 커리큘럼 핵심 주제 영역 표 추가
#    위치: "▶ 과정별 실습 AI 툴 활용 계획" 표 위 (또는 캠퍼스 상세 표 다음)
# ====================================================================
# AI 툴 표를 찾기
tool_table = None
for t in doc.tables:
    if len(t.rows) > 0:
        first = [c.text.strip() for c in t.rows[0].cells]
        if '핵심 실습 툴' in first:
            tool_table = t
            break

if tool_table is not None:
    # AI 툴 표 위쪽 paragraph 찾기 ("▶ 과정별 실습 AI 툴 활용 계획")
    idx = find_para_by_text(doc, "▶ 과정별")
    if idx < 0:
        idx = find_para_by_text(doc, "과정별 실습 AI 툴")

    if idx >= 0:
        # "▶ 과정별..." 단락 바로 위에 새 헤더 + 표 삽입
        # → 그 단락의 이전 element 다음에 삽입
        target_para = doc.paragraphs[idx]
        prev_elem = target_para._element.getprevious()
        if prev_elem is not None:
            anchor = prev_elem
        else:
            anchor = target_para._element  # fallback

        # 5대 주제 영역 표 데이터
        topic_data = [
            ['주제 영역', '세부 내용'],
            ['생성형 AI', '프롬프트 엔지니어링, ChatGPT·Gemini 활용, AI 콘텐츠 생성, AI 기반 리서치'],
            ['데이터 분석', '공공데이터 분석, 시각화, 인사이트 도출, NotebookLM 활용 데이터 해석'],
            ['AI 에이전트', '업무 자동화, 노코드 툴(Zapier·n8n), RAG 기초, Codex 기반 에이전트 설계'],
            ['디지털 콘텐츠 제작', '실감미디어(VR/AR/MR), 영상·이미지 제작 S/W, AI 기반 콘텐츠 생성'],
            ['인문학 × AI 융합', 'AX 시대 인문학의 방향, 문화유산 기반 창작, AI 창작 워크숍'],
        ]

        # 빈 단락 삽입
        insert_paragraph_after(anchor, "", bold=False)
        # 표 삽입
        insert_table_after(anchor, topic_data)
        # 헤더 삽입
        insert_paragraph_after(anchor, "▶ 교육 커리큘럼 핵심 주제 영역", bold=True)
        # 빈 줄 (헤더 위)
        insert_paragraph_after(anchor, "", bold=False)
        print("INSERTED 5대 주제 영역 표")

import os
try:
    doc.save(PATH)
    print(f"\n✅ Saved: {PATH}")
except PermissionError:
    alt_path = '연세AX_제안서_모두의연구소_v2_260506(보완3).docx'
    doc.save(alt_path)
    print(f"\n⚠️ 원본 파일이 잠겨 있어 사본으로 저장: {alt_path}")
    print(f"   Word에서 v2 파일을 닫고 이 사본을 v2로 교체해 주세요.")
