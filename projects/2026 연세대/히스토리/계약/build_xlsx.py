# -*- coding: utf-8 -*-
"""연세 AX 위탁용역 비용 상세 산출 내역서 XLSX 빌더."""
import openpyxl
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "비용상세"

# Styles
thin = Side(border_style="thin", color="000000")
medium = Side(border_style="medium", color="000000")
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
border_thick = Border(left=medium, right=medium, top=medium, bottom=medium)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center", wrap_text=True)
bold = Font(bold=True, name="맑은 고딕", size=10)
bold_lg = Font(bold=True, name="맑은 고딕", size=14)
normal = Font(name="맑은 고딕", size=10)
fill_header = PatternFill("solid", fgColor="D9E1F2")
fill_subtotal = PatternFill("solid", fgColor="FFF2CC")
fill_total = PatternFill("solid", fgColor="FCE4D6")
fill_grandtotal = PatternFill("solid", fgColor="F4B084")
fill_section = PatternFill("solid", fgColor="E2EFDA")

# Column widths
col_widths = [28, 12, 16, 38, 8, 8, 14, 16, 40]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

row = 1

# Title
ws.cell(row=row, column=1, value="연세 AX(AI Transformation) 역량 강화 교육 위탁 용역 — 비용 상세 산출 내역서")
ws.cell(row=row, column=1).font = bold_lg
ws.cell(row=row, column=1).alignment = center
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
ws.row_dimensions[row].height = 28
row += 2

# 공급자 정보
header_info = [
    ("사업명", "연세 AX(AI Transformation) 역량 강화 교육 위탁 용역", "등록번호", "517-88-00184"),
    ("발주처", "학교법인 연세대학교 학술문화처 도서관 디지털미디어서비스팀", "대표이사", "김 승 일 (인)"),
    ("계약 기간", "계약체결일 ~ 2027. 2. 28.", "사업장 소재지", "서울시 강남구 강남대로 324 역삼디오슈페리움 2층, 201·202호"),
    ("제출일자", "2026. 5. 26.", "업태/종목", "서비스 / 경영컨설팅"),
    ("합계 금액", "₩369,600,000 (VAT포함, 一金 삼억육천구백육십만 원)", "담당자", "정윤지 (yj.jeong@modulabs.co.kr)"),
]
for label1, val1, label2, val2 in header_info:
    ws.cell(row=row, column=1, value=label1).font = bold
    ws.cell(row=row, column=1).fill = fill_header
    ws.cell(row=row, column=1).alignment = center
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    ws.cell(row=row, column=2, value=val1).alignment = left
    ws.cell(row=row, column=2).font = normal
    ws.cell(row=row, column=6, value=label2).font = bold
    ws.cell(row=row, column=6).fill = fill_header
    ws.cell(row=row, column=6).alignment = center
    ws.merge_cells(start_row=row, start_column=7, end_row=row, end_column=9)
    ws.cell(row=row, column=7, value=val2).alignment = left
    ws.cell(row=row, column=7).font = normal
    for c in range(1, 10):
        ws.cell(row=row, column=c).border = border_all
    ws.row_dimensions[row].height = 22
    row += 1

row += 1

# 총계 요약 섹션
ws.cell(row=row, column=1, value="■ 총계 요약 (가격제안서 산식)").font = bold_lg
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
ws.cell(row=row, column=1).fill = fill_section
row += 1

summary_headers = ["비목", "구분", "", "", "", "", "산식", "금액(원)", "구성비"]
for c, h in enumerate(summary_headers, 1):
    cell = ws.cell(row=row, column=c, value=h)
    cell.font = bold
    cell.alignment = center
    cell.fill = fill_header
    cell.border = border_all
row += 1

summary_rows = [
    ("인건비(A)", "특급기술자 2명", "", 2_000_000, "1%"),
    ("", "고급기술자 4명", "", 24_000_000, "6%"),
    ("", "중급기술자 4명", "", 16_800_000, "5%"),
    ("", "인건비 합계 (A)", "", 42_800_000, "12%", True),
    ("직접경비(B)", "신촌캠퍼스 도서관 — 연세 AX 아카데미", "", 150_000_000, "41%"),
    ("", "국제캠퍼스 도서관 — 연세 AX 스타트", "", 75_000_000, "20%"),
    ("", "신촌캠퍼스 박물관 — 연세 AX 뮤즈로그", "", 25_000_000, "7%"),
    ("", "직접경비 합계 (B)", "", 250_000_000, "68%", True),
    ("누계 (C)", "A + B", "A + B", 292_800_000, "79%"),
    ("일반관리비 (D)", "(A + B) × 5%", "(A + B) × 5%", 14_640_000, "4%"),
    ("이윤 (E)", "(C + D) × 10%", "(C + D) × 10%", 30_744_000, "8%"),
    ("총계 (F)", "C + D + E", "C + D + E", 338_184_000, "92%", True),
    ("부가세 (G)", "F × 10%", "F × 10%", 33_818_400, "9%"),
    ("교육기부 (H)", "차감 (사회공헌·천원단위 절사)", "차감", -2_402_400, "—"),
    ("합 계", "F + G + H", "F + G + H", 369_600_000, "—", "grand"),
]
for r in summary_rows:
    label = r[0]
    desc = r[1]
    formula = r[2]
    amount = r[3]
    ratio = r[4]
    flag = r[5] if len(r) > 5 else None

    ws.cell(row=row, column=1, value=label).font = bold
    ws.cell(row=row, column=1).alignment = center
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
    ws.cell(row=row, column=2, value=desc).alignment = left
    ws.cell(row=row, column=2).font = bold if (flag == True or flag == "grand") else normal
    ws.cell(row=row, column=7, value=formula).alignment = center
    ws.cell(row=row, column=7).font = normal
    ws.cell(row=row, column=8, value=amount).alignment = right
    ws.cell(row=row, column=8).number_format = "#,##0"
    ws.cell(row=row, column=8).font = bold if (flag == True or flag == "grand") else normal
    ws.cell(row=row, column=9, value=ratio).alignment = center
    ws.cell(row=row, column=9).font = normal

    if flag == "grand":
        for c in range(1, 10):
            ws.cell(row=row, column=c).fill = fill_grandtotal
            ws.cell(row=row, column=c).font = bold
    elif flag == True:
        for c in range(1, 10):
            ws.cell(row=row, column=c).fill = fill_subtotal
    for c in range(1, 10):
        ws.cell(row=row, column=c).border = border_all
    ws.row_dimensions[row].height = 20
    row += 1

row += 1

# 사업별 세부 산출 내역 헤더
ws.cell(row=row, column=1, value="■ 사업별 세부 산출 내역").font = bold_lg
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
ws.cell(row=row, column=1).fill = fill_section
row += 1

detail_headers = ["항목", "비목", "세목", "세부 내역", "인원", "시간", "단가(원)", "공급가액(원)", "비고"]
for c, h in enumerate(detail_headers, 1):
    cell = ws.cell(row=row, column=c, value=h)
    cell.font = bold
    cell.alignment = center
    cell.fill = fill_header
    cell.border = border_all
ws.row_dimensions[row].height = 24
header_row = row
row += 1

# Data: (항목, 비목, 세목, 세부내역, 인원, 시간, 단가, 공급가액, 비고)
def add_section_header(title, fill=None):
    global row
    ws.cell(row=row, column=1, value=title).font = bold
    ws.cell(row=row, column=1).alignment = left
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    ws.cell(row=row, column=1).fill = fill or fill_section
    for c in range(1, 10):
        ws.cell(row=row, column=c).border = border_all
    ws.row_dimensions[row].height = 22
    row += 1

def add_row(data, subtotal=False):
    global row
    for c, v in enumerate(data, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.border = border_all
        cell.font = bold if subtotal else normal
        if c in (5, 6, 7, 8):
            cell.alignment = right
            if c in (7, 8) and isinstance(v, (int, float)):
                cell.number_format = "#,##0"
        elif c in (1, 2, 3):
            cell.alignment = center
        else:
            cell.alignment = left
        if subtotal:
            cell.fill = fill_subtotal
    ws.row_dimensions[row].height = 20
    row += 1

# A. 인건비
add_section_header("A. 인건비 (₩42,800,000)")
add_row(["인건비", "인건비", "특급기술자", "PO·교육철학 자문", 2, 1, 1_000_000, 2_000_000, "사업 총괄 의사결정·교육철학 수립·예산 최종 승인 (자문성 단기 투입)"])
add_row(["인건비", "인건비", "고급기술자", "PM·교육기획 리드·콘텐츠개발 총괄·국제캠퍼스 기획", 4, 1, 6_000_000, 24_000_000, "전체 일정·산출물 관리, 커리큘럼 설계 총괄, 콘텐츠 개발 관리, 캠퍼스별 운영 기획 (2026.05~2027.02 상시)"])
add_row(["인건비", "인건비", "중급기술자", "교육기획·운영·콘텐츠 개발·홍보 마케팅", 4, 1, 4_200_000, 16_800_000, "신촌·국제·박물관 프로그램 실무 기획·운영, 교재 개발, 모집·홍보, 만족도 데이터 관리 (2026.05~2027.02 상시)"])
add_row(["소계 (A)", "", "", "", "", "", "", 42_800_000, ""], subtotal=True)

# B-1. 신촌 아카데미
add_section_header("B-1. 직접경비 — 신촌캠퍼스 도서관 / 연세 AX 아카데미 (₩150,000,000)")
sinchon = [
    ("신촌1. 봄학기 디자인씽킹 워크숍(문제정의)", "직접경비", "강사비", "디자인씽킹 전문 강사", 1, 12, 500_000, 6_000_000, "4회 × 3H, 회당 50명 기준"),
    ("", "직접경비", "운영비", "워크북·교재·다과·진행 운영", 1, 1, 2_000_000, 2_000_000, "4회 통합 운영 자재"),
    ("신촌2. 여름 LLM WIKI 워크숍", "직접경비", "강사비", "주강사 (학습·연구를 위한 생성형 AI 활용)", 1, 8, 500_000, 4_000_000, "2회 × 4H, 회당 50명 기준"),
    ("", "직접경비", "운영비", "실습 교재·교구·다과·진행 운영", 1, 1, 1_500_000, 1_500_000, "2회 통합 운영 자재"),
    ("신촌3. 여름 데이터분석 기초 워크숍", "직접경비", "강사비", "주강사 (바이브코딩 데이터분석)", 1, 8, 500_000, 4_000_000, "2회 × 4H, 회당 50명 기준"),
    ("", "직접경비", "운영비", "실습 교재·교구·다과·진행 운영", 1, 1, 1_500_000, 1_500_000, "2회 통합 운영 자재"),
    ("신촌4. 여름 디지털콘텐츠 제작 워크숍", "직접경비", "강사비", "주강사 (AI 활용 디지털콘텐츠)", 1, 8, 500_000, 4_000_000, "2회 × 4H, 회당 50명 기준"),
    ("", "직접경비", "운영비", "실습 교재·교구·다과·진행 운영", 1, 1, 1_500_000, 1_500_000, "2회 통합 운영 자재"),
    ("신촌5. 몰입형 AX 캠프 트랙1 — Agentic AI", "직접경비", "강사비", "주강사 (프롬프트→컨텍스트→하네스)", 1, 40, 500_000, 20_000_000, "2회 × 20H, 회당 60명 기준"),
    ("", "직접경비", "보조강사비", "보조강사 (실습·튜터링)", 1, 40, 200_000, 8_000_000, "트랙 전 회차 보조"),
    ("", "직접경비", "운영비", "식대·다과·교재·AI 인프라·홍보", 1, 1, 2_000_000, 2_000_000, "식대 15,000원/식·AI 공용 계정 70,000원/팀·월 포함"),
    ("신촌6. 몰입형 AX 캠프 트랙2 — 바이브코딩 풀스택", "직접경비", "강사비", "주강사 (바이브코딩 풀스택 웹앱)", 1, 40, 500_000, 20_000_000, "2회 × 20H, 회당 60명 기준"),
    ("", "직접경비", "보조강사비", "보조강사 (실습·튜터링)", 1, 40, 200_000, 8_000_000, "트랙 전 회차 보조"),
    ("", "직접경비", "운영비", "식대·다과·교재·AI 인프라·홍보", 1, 1, 2_000_000, 2_000_000, "식대 15,000원/식·AI 공용 계정 70,000원/팀·월 포함"),
    ("신촌7. 가을 AX 트렌드 세미나", "직접경비", "강사비", "산업 전문가 연사", 4, 2, 1_000_000, 8_000_000, "IT·미디어 산업 전문가 4인, 회당 2H"),
    ("", "직접경비", "운영비", "행사 운영·다과·홍보", 1, 1, 2_000_000, 2_000_000, "4회 통합 운영비"),
    ("신촌8. 가을 디자인씽킹 워크숍(창의적 해결)", "직접경비", "강사비", "디자인씽킹 전문 강사", 1, 12, 500_000, 6_000_000, "4회 × 3H, 회당 50명 기준"),
    ("", "직접경비", "운영비", "워크북·교재·다과·진행 운영", 1, 1, 2_000_000, 2_000_000, "4회 통합 운영 자재"),
    ("신촌9. 가을 아이디어 공모전", "직접경비", "강사비", "심사위원 섭외", 5, 1, 1_000_000, 5_000_000, "분야별 심사위원 5인"),
    ("", "직접경비", "시상금", "우수 출품작 시상", 1, 1, 3_000_000, 3_000_000, "우수 3팀 시상"),
    ("", "직접경비", "운영비", "공모전 운영·플랫폼·홍보", 1, 1, 6_000_000, 6_000_000, "개방형 200팀 규모 운영"),
    ("신촌10. 겨울 바이브코딩 해커톤", "직접경비", "멘토비", "1:N 밀착 멘토링", 10, 24, 100_000, 24_000_000, "멘토 10명 × 3회 × 8H, 100팀 참여"),
    ("", "직접경비", "시상금", "우수 팀 시상", 1, 1, 2_000_000, 2_000_000, "우수 팀 시상"),
    ("", "직접경비", "운영비", "식대·AI 인프라·다과·홍보·진행", 1, 1, 4_000_000, 4_000_000, "100팀 규모 운영 자재"),
    ("신촌11. 겨울 프로젝트 발표회(쇼케이스)", "직접경비", "강사비", "분야별 심사위원·연사", 2, 3, 500_000, 3_000_000, "1회 × 3H 행사"),
    ("", "직접경비", "운영비", "행사 운영·다과·SNS 배포", 1, 1, 500_000, 500_000, "발표회 운영 자재"),
]
for r in sinchon:
    add_row(list(r))
add_row(["소계 (B-1)", "", "", "", "", "", "", 150_000_000, ""], subtotal=True)

# B-2. 국제 스타트
add_section_header("B-2. 직접경비 — 국제캠퍼스 도서관 / 연세 AX 스타트 (₩75,000,000)")
international = [
    ("국제1. AX 트렌드 세미나", "직접경비", "강사비", "산업 전문가 연사", 3, 2, 800_000, 4_800_000, "전문가 3인 × 2H 강의"),
    ("", "직접경비", "운영비", "운영·기념품·다과", 1, 1, 700_000, 700_000, "3회 통합 운영 자재"),
    ("국제2. 생성형 AI 원데이 클래스", "직접경비", "강사비", "주강사 (생성형 AI 입문)", 1, 20, 400_000, 8_000_000, "10회 × 2H, 회당 20명 기준"),
    ("", "직접경비", "운영비", "교재·다과·진행 운영", 1, 1, 1_500_000, 1_500_000, "10회 통합 운영 자재"),
    ("국제3. 미디어 제작·활용 워크숍", "직접경비", "강사비", "주강사 (CapCut·PPT·Adobe 등)", 1, 45, 350_000, 15_750_000, "15회 × 3H, 회당 20명 기준"),
    ("", "직접경비", "운영비", "실습 인프라·교재·다과·진행 운영", 1, 1, 2_250_000, 2_250_000, "15회 통합 운영 자재"),
    ("국제4. 생성형 AI 활용 워크숍", "직접경비", "강사비", "주강사 (생성형 AI 심화 활용)", 1, 15, 400_000, 6_000_000, "5회 × 3H"),
    ("", "직접경비", "운영비", "교재·다과·진행 운영", 1, 1, 500_000, 500_000, "5회 통합 운영 자재"),
    ("국제5. 공모전 연계 생성형 AI·콘텐츠 워크숍", "직접경비", "강사비", "주강사 (공모전 출품 지원)", 1, 45, 350_000, 15_750_000, "15회 × 3H, 회당 20명 기준"),
    ("", "직접경비", "운영비", "실습 인프라·교재·다과·진행 운영", 1, 1, 2_250_000, 2_250_000, "15회 통합 운영 자재"),
    ("국제6. 공모전 기술 워크숍(멘토링)", "직접경비", "멘토비", "공모전 출품작 시제품화 지원", 1, 32, 80_000, 2_560_000, "멘토 1명 × 4회 × 8H"),
    ("", "직접경비", "운영비", "다과·진행 운영", 1, 1, 440_000, 440_000, "4회 통합 운영 자재"),
    ("국제7. AX 기술 역량 강화 공모전", "직접경비", "강사비", "심사위원 섭외", 3, 2, 500_000, 3_000_000, "2회 × 심사위원 3인"),
    ("", "직접경비", "시상금·운영비", "시상금·공모전 운영·홍보", 1, 1, 1_500_000, 1_500_000, "2회 통합 시상·운영비"),
    ("국제8. 북클럽 공모전", "직접경비", "강사비", "심사위원 섭외", 3, 1, 500_000, 1_500_000, "심사위원 3인"),
    ("", "직접경비", "시상금·운영비", "시상금·공모전 운영·홍보", 1, 1, 1_000_000, 1_000_000, "1회 시상·운영비"),
    ("국제9. AX 북클럽 미디어 교육·멘토링", "직접경비", "강사비", "주강사 (미디어 교육)", 1, 16, 250_000, 4_000_000, "8회 × 2H 강의"),
    ("", "직접경비", "멘토비", "결과물 시제품화 멘토링", 1, 64, 50_000, 3_200_000, "8회 × 8H 멘토링"),
    ("", "직접경비", "운영비", "교재·다과·진행 운영", 1, 1, 300_000, 300_000, "8회 통합 운영 자재"),
]
for r in international:
    add_row(list(r))
add_row(["소계 (B-2)", "", "", "", "", "", "", 75_000_000, ""], subtotal=True)

# B-3. 박물관 뮤즈로그
add_section_header("B-3. 직접경비 — 신촌캠퍼스 박물관 / 연세 AX 뮤즈로그 (₩25,000,000)")
museum = [
    ("박물관1. 인문학 역량 강화 세미나", "직접경비", "강사비", "인문학 전문 연사", 2, 2, 800_000, 3_200_000, "「AX 시대 인문학의 미래」 / 「한글을 지킨 연세의 거인들」"),
    ("", "직접경비", "운영비", "운영·기념품·다과", 1, 1, 800_000, 800_000, "2회 통합 운영 자재"),
    ("박물관2. AI 실감형 콘텐츠 창작 워크숍", "직접경비", "강사비", "주강사 (박물관 컬렉션 AI 재해석)", 1, 6, 500_000, 3_000_000, "2회 × 3H, 기초·심화"),
    ("", "직접경비", "운영비", "교재·박물관 인프라 연동·다과·진행", 1, 1, 1_500_000, 1_500_000, "2회 통합 운영 자재"),
    ("박물관3. 문화유산 기반 AX 공모전", "직접경비", "강사비", "심사위원 섭외", 3, 1, 500_000, 1_500_000, "심사위원 3인"),
    ("", "직접경비", "시상금·운영비", "시상금·박물관 전시 운영·홍보", 1, 1, 2_000_000, 2_000_000, "우수작 박물관 전시·디지털 아카이브"),
    ("박물관4. VR/AR/MR 실감형 콘텐츠 5종 개발", "직접경비", "콘텐츠 개발비", "시나리오 기획·제작·검수 (Physical AI 접점 포함)", 5, 1, 2_600_000, 13_000_000, "VR/AR/MR + 공간 인식·센서 연동 5종, 발주처 자산 귀속"),
]
for r in museum:
    add_row(list(r))
add_row(["소계 (B-3)", "", "", "", "", "", "", 25_000_000, ""], subtotal=True)

# C·D·E·F·G·H
add_section_header("C·D·E·F·G·H. 일반관리비·이윤·부가세·교육기부")
add_row(["누계 (C)", "사업비 누계", "", "인건비 + 직접경비 (A + B)", "", "", "", 292_800_000, "42,800,000 + 250,000,000"])
add_row(["일반관리비 (D)", "일반관리비", "", "SW사업대가 기준 일반관리비 (A + B) × 5%", "", "", "", 14_640_000, "정보통신산업진흥원 기준 5% 적용"])
add_row(["이윤 (E)", "이윤", "", "SW사업대가 기준 이윤 (C + D) × 10%", "", "", "", 30_744_000, "정보통신산업진흥원 기준 10% 적용"])
add_row(["총계 (F)", "공급가액 합계", "", "C + D + E", "", "", "", 338_184_000, "부가세 별도"], subtotal=True)
add_row(["부가세 (G)", "부가가치세", "", "F × 10%", "", "", "", 33_818_400, ""])
add_row(["교육기부 (H)", "교육기부", "", "차감 (사회공헌·천원단위 절사)", "", "", "", -2_402_400, "한글표기 원단위 절사 및 사회공헌 차원 자율 기부"])

# 합계
ws.cell(row=row, column=1, value="합 계").font = bold
ws.cell(row=row, column=1).alignment = center
ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
ws.cell(row=row, column=2, value="F + G + H (VAT포함, 一金 삼억육천구백육십만 원)").alignment = left
ws.cell(row=row, column=2).font = bold
ws.cell(row=row, column=8, value=369_600_000).alignment = right
ws.cell(row=row, column=8).number_format = "#,##0"
ws.cell(row=row, column=8).font = bold
ws.cell(row=row, column=9, value="VAT포함")
ws.cell(row=row, column=9).alignment = center
ws.cell(row=row, column=9).font = bold
for c in range(1, 10):
    ws.cell(row=row, column=c).fill = fill_grandtotal
    ws.cell(row=row, column=c).border = border_all
ws.row_dimensions[row].height = 24
row += 2

# 산정 기준 (각주)
ws.cell(row=row, column=1, value="■ 산정 기준 (각주)").font = bold_lg
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
ws.cell(row=row, column=1).fill = fill_section
row += 1

notes = [
    "1. 주강사 단가: 모두의연구소 사내 표준 단가표 기준. 산업 전문가 연사(세미나) 800,000~1,000,000원/시간 · 디자인씽킹·AX 캠프·심화 워크숍 주강사 500,000원/시간 · 입문 단계 원데이/미디어 워크숍 주강사 350,000~400,000원/시간.",
    "2. 보조강사·멘토 단가: 보조강사 200,000원/시간 (AX 캠프 실습 보조) · 멘토(해커톤) 100,000원/시간 (1:N 밀착 멘토링) · 멘토(공모전 출품 지원) 50,000~80,000원/시간.",
    "3. 운영비 포함 항목: 교재 인쇄·실습 환경 세팅·식대(15,000원/식)·다과·홍보물·교통/체재비·기타 소모품 일괄 포함.",
    "4. AI 인프라: ChatGPT·Claude 등 AI 도구 공용 계정 70,000원/팀·월 기준 (AX 캠프·해커톤 등 집중 운영 기간 한정).",
    "5. 시상금: 공모전·해커톤 우수 팀에 대한 시상금. RFP 명시 고정금액에 준함.",
    "6. 인건비 등급 정의: 정보통신산업진흥원 「SW사업대가 산정 가이드」의 기술자 등급(특급/고급/중급) 기준. 본 사업 참여인력 10인을 등급별 배치 (PO 1·수행총괄 PM 1·교육기획 4·교육운영 1+보조·콘텐츠개발 3).",
    "7. 콘텐츠 개발비: 실감형 콘텐츠 5종(VR/AR/MR + Physical AI 접점)에 한정. 콘텐츠 PM·기획·시나리오 설계·인프라 연계·검수 일체 포함. 산출물은 발주처 자산으로 귀속.",
    "8. 예산 집행 원칙: 일정·운영 환경에 따라 라인별 변동 가능, 총액 범위 내 집행. 항목 간 전용(轉用) 필요 시 발주처 사유서 제출 → 서면 승인 후 집행.",
    "9. 교육기부: 본 사업의 사회적 가치 실현 및 한글표기 원단위 절사를 위해 부가세 포함 합계 기준 ₩2,402,400을 자율 기부.",
]
for n in notes:
    ws.cell(row=row, column=1, value=n).font = normal
    ws.cell(row=row, column=1).alignment = left
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    ws.row_dimensions[row].height = 28
    row += 1

row += 1
ws.cell(row=row, column=1, value="* 상기 견적은 일정·운영 환경에 따라 변경될 수 있습니다. 문의: 정윤지 PM (yj.jeong@modulabs.co.kr)").font = Font(italic=True, name="맑은 고딕", size=9)
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
ws.cell(row=row, column=1).alignment = left
row += 1

# Freeze pane below header row
ws.freeze_panes = f"A{header_row + 1}"

# Save
out = r"C:\Users\Admin\Downloads\ai-education-proposal\projects\2026 연세대\계약\[모두의연구소]연세AX_위탁용역_비용상세_260526.xlsx"
wb.save(out)
print(f"saved: {out}")
