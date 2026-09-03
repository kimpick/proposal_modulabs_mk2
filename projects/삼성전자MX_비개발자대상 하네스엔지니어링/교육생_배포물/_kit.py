# -*- coding: utf-8 -*-
"""교육생 배포물 docx 공통 스타일 모듈 (모두의연구소 기업교육팀)"""
import re
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

KR = '맑은 고딕'
MONO = 'Consolas'
RED = 'EE3E4C'
RED_DK = 'C92434'
INK = '111111'
BODY = '333333'
MUTED = '5B6270'
FAINT = '8A919C'
BORDER = 'E4E4E8'
LINE = 'F0F0F2'
SOFT = 'FAFAFB'
SOFT2 = 'F2F3F6'
PINK = 'FFF6F7'
PINK_B = 'F4CCD0'


def _rgb(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn('w:' + k), v)
    return e


def new_doc():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin, s.bottom_margin = Mm(16), Mm(15)
    s.left_margin, s.right_margin = Mm(18), Mm(18)
    st = doc.styles['Normal']
    st.font.name = KR
    st.font.size = Pt(10.5)
    st.font.color.rgb = _rgb(BODY)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), KR)
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.42
    pf.space_before, pf.space_after = Pt(0), Pt(4)
    # 한글과 영문·숫자 사이에 Word가 넣는 자동 간격 제거 (배포물 조판 품질)
    ppr = st.element.get_or_add_pPr()
    for tag in ('w:autoSpaceDE', 'w:autoSpaceDN'):
        ppr.append(_el(tag, val='0'))
    return doc


def _set_font(run, name=KR, size=10.5, bold=False, color=BODY, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), KR if name == KR else name)
    return run


def rich(p, text, size=10.5, color=BODY, bold_color=INK, base_bold=False):
    """**굵게**, `코드`, ____ 밑줄 빈칸 지원"""
    for tok in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|_{3,})', text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            # 굵게 안에 들어온 `코드` 도 살려서 렌더 (백틱이 글자로 새는 것 방지)
            for inner in re.split(r'(`[^`]+`)', tok[2:-2]):
                if not inner:
                    continue
                if inner.startswith('`') and inner.endswith('`'):
                    _set_font(p.add_run(inner[1:-1]), name=MONO, size=size - 0.5,
                              bold=True, color=RED_DK)
                else:
                    _set_font(p.add_run(inner), size=size, bold=True, color=bold_color)
        elif tok.startswith('`') and tok.endswith('`'):
            _set_font(p.add_run(tok[1:-1]), name=MONO, size=size - 0.5, color=RED_DK)
        elif set(tok) == {'_'}:
            r = _set_font(p.add_run(' ' * (len(tok) * 2)), size=size, color=FAINT)
            r.font.underline = True
        else:
            _set_font(p.add_run(tok), size=size, bold=base_bold, color=color)
    return p


def _borders(el, edges, sz=4, color=BORDER, tag='w:pBdr'):
    pr = el.get_or_add_pPr() if tag == 'w:pBdr' else el
    old = pr.find(qn(tag))
    if old is not None:
        pr.remove(old)
    bd = OxmlElement(tag)
    for e, spec in edges.items():
        bd.append(_el('w:' + e, val=spec.get('val', 'single'), sz=str(spec.get('sz', sz)),
                      space=str(spec.get('space', 0)), color=spec.get('color', color)))
    pr.append(bd)


def _row_opts(row, cant_split=True, header=False):
    trPr = row._tr.get_or_add_trPr()
    if cant_split:
        trPr.append(OxmlElement('w:cantSplit'))
    if header:
        trPr.append(OxmlElement('w:tblHeader'))
    return row


def shade(cell, fill):
    cell._tc.get_or_add_tcPr().append(_el('w:shd', val='clear', color='auto', fill=fill))


def doc_frame(doc, header_text, footer_text):
    """공식 문서 조판 — 머리말 한 줄과 쪽번호(현재/전체)"""
    sec = doc.sections[0]
    sec.header_distance, sec.footer_distance = Mm(10), Mm(10)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_font(hp.add_run(header_text), size=8.5, color=FAINT)
    _borders(hp._p, {'bottom': {'val': 'single', 'sz': 4, 'space': 4, 'color': BORDER}})

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(fp.add_run(footer_text + '     '), size=8.5, color=FAINT)
    _field(fp, 'PAGE')
    _set_font(fp.add_run(' / '), size=8.5, color=FAINT)
    _field(fp, 'NUMPAGES')
    return doc


def _field(p, instr):
    r1 = p.add_run()
    _set_font(r1, size=8.5, color=FAINT)
    r1._element.append(_el('w:fldChar', fldCharType='begin'))
    r2 = p.add_run()
    _set_font(r2, size=8.5, color=FAINT)
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = ' %s ' % instr
    r2._element.append(it)
    r3 = p.add_run()
    _set_font(r3, size=8.5, color=FAINT)
    r3._element.append(_el('w:fldChar', fldCharType='end'))
    return p


def title_block(doc, title, sub, meta=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    _set_font(p.add_run(title), size=19, bold=True, color=INK)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(7)
    _set_font(p2.add_run(sub), size=10, color=MUTED)
    _borders(p2._p, {'bottom': {'val': 'single', 'sz': 20, 'space': 6, 'color': RED}})
    if meta:
        p3 = doc.add_paragraph()
        p3.paragraph_format.space_before = Pt(5)
        p3.paragraph_format.space_after = Pt(9)
        _set_font(p3.add_run(meta), size=9.5, color=FAINT)
    return doc


def h2(doc, num, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(6)
    if num:
        _set_font(p.add_run(num + '  '), size=13, bold=True, color=RED)
    _set_font(p.add_run(text), size=13, bold=True, color=INK)
    p.paragraph_format.keep_with_next = True
    _borders(p._p, {'bottom': {'val': 'single', 'sz': 6, 'space': 3, 'color': BORDER}})
    return p


def h3(doc, text, note=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(4)
    _set_font(p.add_run(text), size=11, bold=True, color=INK)
    if note:
        _set_font(p.add_run('   ' + note), size=9.5, color=FAINT)
    p.paragraph_format.keep_with_next = True
    return p


def para(doc, text, size=10.5, color=BODY, before=0, after=4, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if indent:
        p.paragraph_format.left_indent = Mm(indent)
    rich(p, text, size=size, color=color)
    return p


def bullets(doc, items, size=10.5, marker='·', indent=3.5, hang=3.5):
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3.5)
        p.paragraph_format.left_indent = Mm(indent + hang)
        p.paragraph_format.first_line_indent = Mm(-hang)
        _set_font(p.add_run(marker + '  '), size=size, color=RED)
        rich(p, it, size=size)
    return doc


def checklist(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4.5)
        p.paragraph_format.left_indent = Mm(7)
        p.paragraph_format.first_line_indent = Mm(-7)
        _set_font(p.add_run('☐   '), size=size + 1, color=MUTED)
        rich(p, it, size=size)
    return doc


def _grid(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = float(sum(widths))
    avail = 174.0
    for row in table.rows:
        for i, c in enumerate(row.cells):
            c.width = Mm(widths[i] / total * avail)


def table(doc, headers, rows, widths, size=10, head_fill=INK, zebra=True,
          key_col=0, head_size=9.5, row_pad=True):
    t = doc.add_table(rows=0, cols=len(widths))
    t.style = 'Table Grid'
    _borders(t._tbl.tblPr, {
        'top': {'sz': 4, 'color': BORDER}, 'left': {'val': 'none', 'sz': 0},
        'bottom': {'sz': 4, 'color': BORDER}, 'right': {'val': 'none', 'sz': 0},
        'insideH': {'sz': 4, 'color': LINE}, 'insideV': {'val': 'none', 'sz': 0},
    }, tag='w:tblBorders')
    if headers:
        r = t.add_row()
        _row_opts(r, header=True)
        for i, htxt in enumerate(headers):
            c = r.cells[i]
            shade(c, head_fill)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            _set_font(p.add_run(htxt), size=head_size, bold=True, color='FFFFFF')
    for n, row in enumerate(rows):
        r = t.add_row()
        _row_opts(r)
        for i, val in enumerate(row):
            c = r.cells[i]
            if zebra and n % 2 == 1:
                shade(c, SOFT)
            lines = val.split('\n') if isinstance(val, str) else [str(val)]
            for j, ln in enumerate(lines):
                p = c.paragraphs[0] if j == 0 else c.add_paragraph()
                p.paragraph_format.space_before = Pt(4 if (j == 0 and row_pad) else 0.5)
                p.paragraph_format.space_after = Pt(4 if (j == len(lines) - 1 and row_pad) else 0.5)
                p.paragraph_format.line_spacing = 1.34
                bold = (i == key_col and j == 0)
                rich(p, ln, size=size, color=INK if bold else BODY, base_bold=bold)
    _grid(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def formtable(doc, rows, label_w=42, size=10, height=9):
    """빈칸 채우기 표 — 왼쪽 라벨(음영) / 오른쪽 입력칸"""
    t = doc.add_table(rows=0, cols=2)
    t.style = 'Table Grid'
    _borders(t._tbl.tblPr, {
        'top': {'sz': 4, 'color': BORDER}, 'left': {'sz': 4, 'color': BORDER},
        'bottom': {'sz': 4, 'color': BORDER}, 'right': {'sz': 4, 'color': BORDER},
        'insideH': {'sz': 4, 'color': BORDER}, 'insideV': {'sz': 4, 'color': BORDER},
    }, tag='w:tblBorders')
    for label, hint, h in rows:
        r = t.add_row()
        _row_opts(r)
        r.height = Mm(h if h else height)
        c0, c1 = r.cells
        shade(c0, SOFT2)
        p = c0.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.34
        rich(p, label, size=size, color=INK, base_bold=True)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        if hint:
            _set_font(p1.add_run(hint), size=size - 0.5, color=FAINT)
    _grid(t, [label_w, 100 - label_w])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def box(doc, lines, fill=SOFT, border=BORDER, size=10.5, accent=None, mono=False):
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    shade(c, fill)
    _borders(t._tbl.tblPr, {
        'top': {'sz': 4, 'color': border}, 'bottom': {'sz': 4, 'color': border},
        'left': {'sz': 18 if accent else 4, 'color': accent or border},
        'right': {'sz': 4, 'color': border},
        'insideH': {'val': 'none', 'sz': 0}, 'insideV': {'val': 'none', 'sz': 0},
    }, tag='w:tblBorders')
    for j, ln in enumerate(lines):
        p = c.paragraphs[0] if j == 0 else c.add_paragraph()
        p.paragraph_format.space_before = Pt(5 if j == 0 else 2)
        p.paragraph_format.space_after = Pt(5 if j == len(lines) - 1 else 2)
        p.paragraph_format.line_spacing = 1.4
        if mono:
            _set_font(p.add_run(ln), name=MONO, size=size, color=INK)
        else:
            rich(p, ln, size=size)
    _grid(t, [100])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def prompt_box(doc, lines, caption=None):
    if caption:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        _set_font(p.add_run(caption), size=9.5, bold=True, color=RED_DK)
    return box(doc, lines, fill=SOFT2, border=BORDER, size=10, accent=RED, mono=True)


def footer_note(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(3)
    _borders(p._p, {'top': {'val': 'dashed', 'sz': 4, 'space': 6, 'color': BORDER}})
    for i, ln in enumerate(lines):
        q = p if i == 0 else doc.add_paragraph()
        q.paragraph_format.space_after = Pt(2)
        rich(q, ln, size=9, color=FAINT)
    return doc


def brand_footer(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(p.add_run(text), size=8.5, color=FAINT)
    return doc


def page_break(doc):
    doc.add_page_break()
    return doc
