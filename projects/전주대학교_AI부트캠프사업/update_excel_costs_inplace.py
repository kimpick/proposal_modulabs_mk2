import openpyxl

file_path = r"c:\Users\Admin\Downloads\ai-education-proposal\projects\전주대학교_AI부트캠프사업\견적\[비즈팀] 전주대_AI 부트캠프_예산산출.xlsx"
out_path = r"c:\Users\Admin\Downloads\ai-education-proposal\projects\전주대학교_AI부트캠프사업\견적\[비즈팀] 전주대_AI 부트캠프_예산산출.xlsx"
try:
    wb = openpyxl.load_workbook(file_path)
    sheets_to_update = ['ND1', 'ND2', 'ND5']
    for sheet_name in sheets_to_update:
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=10, max_row=40):
                if row[2].value == 'AI 툴 사용료':
                    row[3].value = 40000 
                    row[4].value = 30
                    row[5].value = "명"
                    row[6].value = 1
                    row[7].value = "식"
                    row[8].value = 1
                    row[9].value = "회"
                    row[10].value = "=D{}*E{}*G{}*I{}".format(row[0].row, row[0].row, row[0].row, row[0].row)
                    row[11].value = "OpenAI, Claude 등 API 활용 크레딧 (인당 4만원, 모듈 공통실습)"
                
                if row[2].value == '클라우드 사용료':
                    row[3].value = 2500 
                    row[4].value = 30
                    row[5].value = "명"
                    row[6].value = 60 # 60시간 기준(과정당 4개 모듈 합산)
                    row[7].value = "시간"
                    row[8].value = 1
                    row[9].value = "회"
                    row[10].value = "=D{}*E{}*G{}*I{}".format(row[0].row, row[0].row, row[0].row, row[0].row)
                    row[11].value = "카카오 클라우드 고성능 GPU(분산학습) 및 보조 클러스터 시간당 과금"

                if row[2].value == '기타 구독료':
                    row[3].value = 30000
                    row[4].value = 30
                    row[5].value = "명"
                    row[6].value = 1 # 보통 1시즌 구독
                    row[7].value = "건"
                    row[8].value = 1
                    row[9].value = "식"
                    row[10].value = "=D{}*E{}*G{}*I{}".format(row[0].row, row[0].row, row[0].row, row[0].row)
                    row[11].value = "AI Augmented Coding (자동채점, Web-IDE 등) 구동용"

    wb.save(out_path)
    print("Success overwriting original")
except Exception as e:
    print("Error:", e)
