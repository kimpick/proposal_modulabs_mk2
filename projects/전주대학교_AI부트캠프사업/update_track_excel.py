import openpyxl

file_path = r"c:\Users\Admin\Downloads\ai-education-proposal\projects\전주대학교_AI부트캠프사업\견적\[비즈팀] 전주대_AI 부트캠프_예산산출.xlsx"

try:
    wb = openpyxl.load_workbook(file_path)
    
    # Track configurations: (클라우드 단가, API 단가, API 비고, 클라우드 비고)
    track_configs = {
        'ND1': {
            'cloud_price': 1000, 
            'api_price': 10000,
            'api_desc': "기초 API 실습 크레딧 (인당 1만원)",
            'cloud_desc': "기본 실습 Web-IDE 및 경량 인스턴스 (시간당 1,000원)"
        },
        'ND2': {
            'cloud_price': 4000,
            'api_price': 15000,
            'api_desc': "기본 API 및 외부 데이터셋 호출 (인당 1.5만원)",
            'cloud_desc': "고성능 GPU VDI (Isaac Sim 등) 및 스토리지 (시간당 4,000원)"
        },
        'ND5': {
            'cloud_price': 2500,
            'api_price': 60000,
            'api_desc': "대규모 상용 LLM API 리소스 (RAG, 에이전트 실험용 / 인당 6만원)",
            'cloud_desc': "AI 서빙 클러스터 및 중급 GPU 인스턴스 (시간당 2,500원)"
        }
    }

    for sheet_name, config in track_configs.items():
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=10, max_row=40):
                # 공통 로직: 수량, 단위 등
                
                if row[2].value == 'AI 툴 사용료':
                    row[3].value = config['api_price'] 
                    row[4].value = 30
                    row[5].value = "명"
                    row[6].value = 1
                    row[7].value = "식"
                    row[8].value = 1
                    row[9].value = "회"
                    row[10].value = "=D{}*E{}*G{}*I{}".format(row[0].row, row[0].row, row[0].row, row[0].row)
                    row[11].value = config['api_desc']
                
                if row[2].value == '클라우드 사용료':
                    row[3].value = config['cloud_price'] 
                    row[4].value = 30
                    row[5].value = "명"
                    row[6].value = 60 # 60시간
                    row[7].value = "시간"
                    row[8].value = 1
                    row[9].value = "회"
                    row[10].value = "=D{}*E{}*G{}*I{}".format(row[0].row, row[0].row, row[0].row, row[0].row)
                    row[11].value = config['cloud_desc']

                if row[2].value == '기타 구독료':
                    row[3].value = 30000
                    row[4].value = 30
                    row[5].value = "명"
                    row[6].value = 1
                    row[7].value = "건"
                    row[8].value = 1
                    row[9].value = "식"
                    row[10].value = "=D{}*E{}*G{}*I{}".format(row[0].row, row[0].row, row[0].row, row[0].row)
                    row[11].value = "AI Augmented Coding (자동채점 등) 구독료 (인당 3만원)"

    wb.save(file_path)
    print("Successfully updated Excel with customized track configurations.")
except Exception as e:
    print("Error:", e)
