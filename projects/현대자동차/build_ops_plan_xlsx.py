from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# 기존 현대자동차 엑셀 스크립트와 동일 톤의 스타일
HEADER_FILL = PatternFill("solid", start_color="1F4E78")
HEADER_FONT = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="맑은 고딕", size=10)
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws: Any, row: int, col_count: int) -> None:
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def style_body(ws: Any, start_row: int, end_row: int, col_count: int) -> None:
    for r in range(start_row, end_row + 1):
        for c in range(1, col_count + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = BORDER


def set_widths(ws: Any, widths: list[int]) -> None:
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def parse_hours(text: str) -> float:
    if not text:
        return 0.0
    match = re.search(r"(\d+(?:\.\d+)?)\s*H", text, flags=re.IGNORECASE)
    if match:
        return float(match.group(1))
    match = re.search(r"(\d+(?:\.\d+)?)\s*시간", text)
    if match:
        return float(match.group(1))
    return 0.0


def clean_item(item: str) -> str:
    if not item:
        return ""
    return re.sub(r"^\[\d+(?:\.\d+)?H\]\s*", "", item).strip()


def parse_requirements_feedback(req_md: Path) -> list[str]:
    if not req_md.exists():
        return []
    lines = req_md.read_text(encoding="utf-8").splitlines()
    in_feedback = False
    feedback: list[str] = []
    for line in lines:
        if line.startswith("## 4."):
            in_feedback = True
            continue
        if in_feedback and line.startswith("## "):
            break
        if in_feedback and line.strip().startswith("- "):
            feedback.append(line.strip()[2:].strip())
    return feedback


def infer_day_input(day_index: int) -> str:
    if day_index <= 5:
        return "MS Office/Confluence 데이터, 실습용 코드 베이스, 전일 학습 산출물"
    if day_index == 6:
        return "2주차 개별 현업 미션 산출물, 사내 보안 가이드라인(CLAUDE.md), 코드 리포지토리"
    return "보안 테스트베드 로그, CVE/KEV/EPSS 데이터, 전일 프로젝트 산출물"


def infer_day_output(module_name: str) -> str:
    mapping = [
        ("기초", "기초 실습 노트 및 초기 코드"),
        ("파이프라인", "모듈별 파이프라인 코드/설계 결과"),
        ("고도화", "검색 품질 개선 결과 및 평가 로그"),
        ("메모리", "Memory/HITL 설계 문서"),
        ("Q&A", "팀별 질의응답 반영 개선안"),
        ("미션 리뷰", "미션 리뷰 리포트 및 개선 백로그"),
        ("실행 자동화", "보안 자동화 스크립트 및 로그 파서"),
        ("중앙 통제소", "라우팅/승인 통제 플로우"),
        ("Threat Intelligence", "취약점 시나리오 분석 결과"),
        ("통합 PoC", "최종 통합 PoC 데모 및 운영 메모"),
    ]
    for key, value in mapping:
        if key in module_name:
            return value
    return "당일 실습 결과물 및 팀별 공유 자료"


def infer_operational_point(module_name: str, activities: str) -> str:
    text = f"{module_name} {activities}"
    if "Multi Query" in text:
        return "Multi Query/Re-Ranking 적용 전후 품질 차이를 정량 지표로 비교"
    if "CVE/KEV/EPSS" in text:
        return "TI 데이터 최신성 체크와 내부 정책 충돌 여부를 사전 점검"
    if "HITL" in text:
        return "위험 명령은 승인 게이트를 통과하도록 운영 룰을 고정"
    if "미션 리뷰" in text:
        return "미션 리뷰 세션을 Day5 Q&A와 분리 운영해 피드백 누락 방지"
    return "강의-실습-리뷰 전환 시점마다 이해도 체크 및 즉시 보정"


def build_overview_sheet(ws: Any, data: dict[str, Any], feedback_lines: list[str]) -> None:
    headers = ["항목", "내용", "비고"]
    ws.append(headers)
    style_header(ws, 1, len(headers))

    feedback_summary = "\n".join(f"- {line}" for line in feedback_lines) if feedback_lines else "요구조건.md 기준 상세 피드백 미기재"

    rows = [
        ("프로젝트 명", data.get("title", "TBD"), ""),
        ("대상 조직", data.get("info", {}).get("target", "TBD"), ""),
        ("교육 목적", "Agentic RAG + Autonomous Security Agent 현업 내재화", "요구조건.md 기준"),
        ("총 교육시간", data.get("info", {}).get("duration", "총 70시간"), "1차 35H + 미션 + 2차 35H"),
        ("운영 구조", "Day 기준 10일 운영 (Day1~Day10)", "시간 해상도: Day"),
        ("핵심 기술 스택", ", ".join(data.get("tech_stack", [])), "실습/운영 공통"),
        ("핵심 피드백 반영사항", feedback_summary, "요구조건.md §4"),
        ("운영 용어 구분", "2주차 Day1 미션 리뷰와 Day5 내부 프로젝트 Q&A를 분리 운영", "고객 요청 반영"),
        ("운영 산출물", "운영계획서, Day별 진행표, RACI, 체크리스트, 리스크 대응표, KPI 리포팅 체계", "본 워크북 7탭"),
    ]
    for row in rows:
        ws.append(row)

    style_body(ws, 2, len(rows) + 1, len(headers))
    set_widths(ws, [22, 96, 26])
    for r in range(2, len(rows) + 2):
        ws.row_dimensions[r].height = 54
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def build_day_plan_rows(data: dict[str, Any]) -> list[tuple[str, ...]]:
    rows: list[tuple[str, ...]] = []
    day_index = 1
    for track_idx, track in enumerate(data.get("tracks", []), start=1):
        week_block = "1차 교육(35H)" if track_idx == 1 else "2차 교육(35H)"
        for module in track.get("modules", []):
            items = module.get("items", [])
            clean_items = [clean_item(x) for x in items if clean_item(x)]
            time_chunks = []
            for item in items:
                match = re.search(r"\[(\d+(?:\.\d+)?)H\]", item)
                if match:
                    time_chunks.append(f"{match.group(1)}H")
            time_distribution = " / ".join(time_chunks) if time_chunks else module.get("duration", "7H")
            core_goal = clean_items[0] if clean_items else module.get("module_name", "")
            activities = "\n".join(f"- {x}" for x in clean_items) if clean_items else "-"
            day_output = infer_day_output(module.get("module_name", ""))
            op_point = infer_operational_point(module.get("module_name", ""), activities)
            roles = "PM(A) / 주강사(R) / 보조강사(C) / 고객POC(C) / IT지원(I)"

            rows.append(
                (
                    week_block,
                    f"Day{day_index}",
                    module.get("module_name", f"Day{day_index}"),
                    core_goal,
                    time_distribution,
                    activities,
                    infer_day_input(day_index),
                    day_output,
                    roles,
                    op_point,
                )
            )
            day_index += 1
    return rows


def build_day_plan_sheet(ws: Any, data: dict[str, Any]) -> None:
    headers = ["주차블록", "Day", "교육명", "핵심목표", "시간배분", "주요활동", "입력물", "당일산출물", "담당역할", "운영포인트"]
    ws.append(headers)
    style_header(ws, 1, len(headers))

    rows = build_day_plan_rows(data)
    for row in rows:
        ws.append(row)

    style_body(ws, 2, len(rows) + 1, len(headers))
    set_widths(ws, [16, 8, 34, 34, 18, 58, 36, 34, 38, 38])
    for r in range(2, len(rows) + 2):
        ws.row_dimensions[r].height = 86
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def build_module_detail_rows(data: dict[str, Any]) -> list[tuple[str, ...]]:
    table = data.get("curriculum_table", [])
    if table:
        return [
            (
                row.get("course", ""),
                row.get("module_name", ""),
                row.get("module_content", ""),
                row.get("duration", ""),
            )
            for row in table
        ]

    rows: list[tuple[str, ...]] = []
    for track in data.get("tracks", []):
        for module in track.get("modules", []):
            items = module.get("items", [])
            if items:
                for item in items:
                    rows.append(
                        (
                            module.get("module_name", ""),
                            clean_item(item),
                            clean_item(item),
                            module.get("duration", ""),
                        )
                    )
            else:
                rows.append(
                    (
                        module.get("module_name", ""),
                        module.get("module_name", ""),
                        "",
                        module.get("duration", ""),
                    )
                )
    return rows


def build_module_detail_sheet(ws: Any, data: dict[str, Any]) -> None:
    headers = ["교육 명", "모듈 명", "모듈 내용", "소요 시간"]
    ws.append(headers)
    style_header(ws, 1, len(headers))

    rows = build_module_detail_rows(data)
    for row in rows:
        ws.append(row)

    style_body(ws, 2, len(rows) + 1, len(headers))
    set_widths(ws, [34, 34, 88, 14])
    for r in range(2, len(rows) + 2):
        ws.row_dimensions[r].height = 50
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def build_raci_sheet(ws: Any) -> None:
    headers = ["운영업무", "PM", "주강사", "보조강사", "고객POC", "IT지원", "완료기준"]
    ws.append(headers)
    style_header(ws, 1, len(headers))

    rows = [
        ("교육 목표/범위 확정", "A", "R", "C", "C", "I", "목표·범위·우선순위 승인 완료"),
        ("실습 데이터셋 확정", "A", "R", "C", "R", "C", "MS Office/Confluence 샘플 데이터 준비 완료"),
        ("실습 환경 및 계정 점검", "A", "C", "C", "I", "R", "API/네트워크/샌드박스 사전 점검 완료"),
        ("Day1~Day5 강의 운영(1차)", "A", "R", "R", "C", "I", "출석·실습 진행률·질문 처리 기록 완료"),
        ("중간 현업 미션 운영", "A", "C", "R", "R", "I", "미션 제출물/리뷰 일정 확정 및 공유"),
        ("Day6~Day10 강의 운영(2차)", "A", "R", "R", "C", "I", "미션 리뷰와 Q&A 분리 운영 완료"),
        ("보안/컴플라이언스 통제", "A", "R", "C", "C", "R", "HITL 승인 룰 및 로그 정책 적용"),
        ("최종 발표·리포트 취합", "A", "R", "R", "C", "I", "팀별 산출물 및 평가 리포트 취합 완료"),
        ("운영 이관 및 사후 개선", "A", "R", "C", "R", "C", "운영 가이드/개선 백로그 전달 완료"),
    ]
    for row in rows:
        ws.append(row)

    style_body(ws, 2, len(rows) + 1, len(headers))
    set_widths(ws, [34, 10, 10, 12, 12, 10, 62])
    for r in range(2, len(rows) + 2):
        ws.row_dimensions[r].height = 36
        for c in range(2, 7):
            ws.cell(row=r, column=c).alignment = CENTER
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def build_checklist_sheet(ws: Any) -> None:
    headers = ["단계", "체크항목", "상세기준", "담당", "마감", "상태", "증빙"]
    ws.append(headers)
    style_header(ws, 1, len(headers))

    rows = [
        ("D-7", "교육 범위 확정", "10일 운영 범위, 핵심 요구사항, Day별 목표 확정", "PM/고객POC", "TBD", "TBD", "승인 메일/회의록"),
        ("D-7", "실습 데이터 범위 확정", "MS Office·Confluence 샘플/가명화 데이터 확정", "고객POC/주강사", "TBD", "TBD", "데이터 목록"),
        ("D-7", "RACI 공지", "운영역할(R/A/C/I) 내부 공유 및 책임자 확인", "PM", "TBD", "TBD", "역할표"),
        ("D-3", "환경 점검", "API 키/네트워크/보안 정책/샌드박스 접근 점검", "IT지원/보조강사", "TBD", "TBD", "체크로그"),
        ("D-3", "실습 리허설", "Day1~Day2 실습 시나리오 사전 리허설 완료", "주강사/보조강사", "TBD", "TBD", "리허설 노트"),
        ("D-3", "평가 기준 공유", "미션/PoC 평가 항목 및 제출 포맷 공유", "PM/주강사", "TBD", "TBD", "안내문"),
        ("D-1", "참가자 안내", "일정표, 준비물, 접속/실습 안내 전달", "PM", "TBD", "TBD", "공지 메일"),
        ("D-1", "운영 자료 최종화", "강의자료, 실습코드, 체크리스트 최종 점검", "주강사", "TBD", "TBD", "버전 태그"),
        ("D-1", "지원 채널 오픈", "질문/장애 대응 채널 및 에스컬레이션 라인 오픈", "PM/IT지원", "TBD", "TBD", "채널 링크"),
        ("D-Day", "출석 및 운영 브리핑", "출석체크, 일정/규칙/보안 유의사항 안내", "PM/보조강사", "TBD", "TBD", "출석부"),
        ("D-Day", "실습 장애 대응", "환경 오류·접속 오류·권한 이슈 실시간 조치", "IT지원/보조강사", "TBD", "TBD", "장애 처리 로그"),
        ("D-Day", "당일 산출물 회수", "팀별 당일 산출물 취합 및 피드백 반영사항 기록", "주강사/보조강사", "TBD", "TBD", "산출물 폴더"),
    ]
    for row in rows:
        ws.append(row)

    style_body(ws, 2, len(rows) + 1, len(headers))
    set_widths(ws, [10, 28, 62, 16, 14, 10, 34])
    for r in range(2, len(rows) + 2):
        ws.row_dimensions[r].height = 44
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def build_risk_sheet(ws: Any) -> None:
    headers = ["리스크ID", "영역", "리스크설명", "가능성", "영향도", "조기징후", "예방조치", "대응조치", "오너"]
    ws.append(headers)
    style_header(ws, 1, len(headers))

    rows = [
        ("R-01", "데이터", "실습 데이터의 민감정보 포함 가능성", "중", "상", "데이터 반입 승인 지연", "가명화 기준 사전 적용, 반입 승인 체크", "민감 항목 마스킹 후 대체 데이터 사용", "고객POC"),
        ("R-02", "보안", "외부 API/도구 접근 정책 충돌", "중", "상", "실습 중 네트워크 차단 발생", "사전 네트워크 화이트리스트 점검", "오프라인 대체 실습 시나리오 전환", "IT지원"),
        ("R-03", "품질", "Multi Query/Re-Ranking 적용 후 품질 개선 미흡", "중", "중", "정량평가 지표 개선폭 미달", "기준 데이터셋 사전 정제 및 베이스라인 확보", "프롬프트/검색전략 튜닝 세션 추가", "주강사"),
        ("R-04", "운영", "미션 리뷰와 Q&A 일정 충돌", "하", "중", "질문/피드백 세션 장시간 지연", "미션 리뷰(독립)와 Q&A 슬롯 사전 분리", "질문 백로그를 Day 종료 후 별도 세션 처리", "PM"),
        ("R-05", "보안", "HITL 승인 프로세스 누락으로 위험 명령 실행", "하", "상", "승인 로그 누락, 우회 실행 발생", "승인 없이는 실행 불가 룰 고정", "즉시 실행 중지 및 승인 체계 재설정", "주강사"),
        ("R-06", "TI", "CVE/KEV/EPSS 데이터 최신성 부족", "중", "중", "시나리오 결과와 실제 최신 이슈 불일치", "수업 전 최신 피드 스냅샷 확보", "수업 중 동기화/대체 지표 적용", "보조강사"),
        ("R-07", "기술", "OpenClaw/샌드박스 환경 이슈로 실습 중단", "중", "상", "설치 실패/권한 오류 다수 발생", "사전 설치 가이드 및 체크리스트 배포", "원격 지원/대체 VM 환경 즉시 전환", "IT지원"),
        ("R-08", "일정", "팀별 산출물 제출 지연", "중", "중", "마감 직전 산출물 미제출", "중간 체크포인트 운영, 진행률 추적", "필수 산출물 최소 기준으로 범위 재조정", "PM"),
    ]
    for row in rows:
        ws.append(row)

    style_body(ws, 2, len(rows) + 1, len(headers))
    set_widths(ws, [10, 12, 34, 10, 10, 30, 34, 34, 14])
    for r in range(2, len(rows) + 2):
        ws.row_dimensions[r].height = 60
        for c in (1, 4, 5, 9):
            ws.cell(row=r, column=c).alignment = CENTER
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def build_reporting_sheet(ws: Any) -> None:
    headers = ["구분", "항목", "정의", "수집방법", "목표기준", "보고주기", "수신자", "비고"]
    ws.append(headers)
    style_header(ws, 1, len(headers))

    rows = [
        ("정기보고", "일일 운영 리포트", "출석/진행률/장애/요청사항 당일 보고", "운영일지 취합", "당일 100% 발행", "매일", "PM, 고객POC", "Day 종료 후 2시간 내"),
        ("정기보고", "주차 요약 리포트", "1차/2차 교육 블록별 학습성과 요약", "일일 리포트 집계", "주차별 1회", "주차 종료", "PM, 고객POC, 주강사", "핵심 이슈 포함"),
        ("정기보고", "미션 리뷰 리포트", "중간 미션 결과, 개선 우선순위, 재작업 항목", "미션 산출물 리뷰", "리뷰 커버리지 100%", "미션 주간", "PM, 고객POC, 주강사", "Q&A와 분리"),
        ("정기보고", "최종 이관 리포트", "PoC 결과, 운영 가이드, 후속 과제", "최종 발표/산출물 취합", "최종본 1회", "교육 종료", "PM, 고객POC, IT지원", "이관 체크리스트 첨부"),
        ("KPI", "Day 운영 완료율", "계획된 Day 세션 대비 완료 세션 비율", "Day 계획 대비 실적 체크", "100%", "매일", "PM", ""),
        ("KPI", "모듈 커버리지", "TSV형 모듈 32행 중 실제 수행 커버리지", "모듈 체크시트", "95% 이상", "주차 종료", "PM, 주강사", "변경사유 기록"),
        ("KPI", "총 시간 준수율", "계획 70시간 대비 실제 운영 시간 준수", "시간로그", "±5% 이내", "주차 종료", "PM, 고객POC", ""),
        ("KPI", "Multi Query 적용률", "관련 세션에서 Multi Query 실습 수행 비율", "세션 체크리스트", "100%", "1차 종료", "주강사, 고객POC", "Day3 핵심"),
        ("KPI", "CVE/KEV/EPSS 반영률", "TI 실습에서 CVE/KEV/EPSS 연계 수행 비율", "실습 결과 점검", "100%", "2차 종료", "주강사, 고객POC", "Day4 핵심"),
        ("KPI", "HITL 승인룰 적용률", "위험 명령 시 HITL 승인 절차 준수 비율", "승인 로그 검토", "100%", "2차 종료", "PM, IT지원", ""),
        ("KPI", "미션 리뷰 완료율", "중간 미션 제출물 대비 리뷰 완료 비율", "리뷰 로그", "100%", "미션 주간", "PM, 고객POC", "분리 세션 운영"),
        ("KPI", "최종 산출물 완성도", "팀별 데모/리포트/이관문서 제출 완성도", "최종 체크리스트", "90점 이상", "교육 종료", "PM, 고객POC, 주강사", "정성·정량 혼합"),
    ]
    for row in rows:
        ws.append(row)

    style_body(ws, 2, len(rows) + 1, len(headers))
    set_widths(ws, [12, 24, 34, 28, 14, 12, 22, 26])
    for r in range(2, len(rows) + 2):
        ws.row_dimensions[r].height = 44
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def calc_total_hours(data: dict[str, Any]) -> float:
    total = 0.0
    for track in data.get("tracks", []):
        for module in track.get("modules", []):
            total += parse_hours(str(module.get("duration", "")))
    return total


def create_workbook(data: dict[str, Any], req_md: Path, out_path: Path) -> None:
    wb = Workbook()

    # 01_운영개요
    ws1 = wb.active
    ws1.title = "01_운영개요"
    feedback_lines = parse_requirements_feedback(req_md)
    build_overview_sheet(ws1, data, feedback_lines)

    # 02_Day운영계획
    ws2 = wb.create_sheet("02_Day운영계획")
    build_day_plan_sheet(ws2, data)

    # 03_모듈상세(TSV형)
    ws3 = wb.create_sheet("03_모듈상세(TSV형)")
    build_module_detail_sheet(ws3, data)

    # 04_RACI_역할책임
    ws4 = wb.create_sheet("04_RACI_역할책임")
    build_raci_sheet(ws4)

    # 05_사전준비체크리스트
    ws5 = wb.create_sheet("05_사전준비체크리스트")
    build_checklist_sheet(ws5)

    # 06_리스크대응
    ws6 = wb.create_sheet("06_리스크대응")
    build_risk_sheet(ws6)

    # 07_보고체계_KPI
    ws7 = wb.create_sheet("07_보고체계_KPI")
    build_reporting_sheet(ws7)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def parse_args() -> argparse.Namespace:
    script_dir = Path(__file__).resolve().parent
    default_input = script_dir / "제안서" / "content.json"
    default_output = script_dir / "현대자동차_상세운영계획안_7탭.xlsx"

    parser = argparse.ArgumentParser(description="현대자동차 상세 운영 계획안 7탭 엑셀 생성기")
    parser.add_argument("--input-json", default=str(default_input), help="입력 content.json 경로")
    parser.add_argument("--output", default=str(default_output), help="출력 xlsx 경로")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    input_json = Path(args.input_json)
    output_path = Path(args.output)
    req_md = input_json.parent.parent / "요구조건.md"

    if not input_json.exists():
        raise FileNotFoundError(f"입력 JSON 파일이 없습니다: {input_json}")

    with input_json.open("r", encoding="utf-8") as f:
        data = json.load(f)

    create_workbook(data, req_md, output_path)

    day_rows = len(build_day_plan_rows(data))
    module_rows = len(build_module_detail_rows(data))
    total_hours = calc_total_hours(data)

    print(f"생성 완료: {output_path}")
    print("시트:", ["01_운영개요", "02_Day운영계획", "03_모듈상세(TSV형)", "04_RACI_역할책임", "05_사전준비체크리스트", "06_리스크대응", "07_보고체계_KPI"])
    print(f"Day 행수: {day_rows}")
    print(f"TSV형 모듈 행수: {module_rows}")
    print(f"총 교육시간(트랙 모듈 합): {total_hours}H")


if __name__ == "__main__":
    main()
