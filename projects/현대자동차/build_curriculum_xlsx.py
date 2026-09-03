from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

HEADER_FILL = PatternFill("solid", start_color="1F4E78")
HEADER_FONT = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="맑은 고딕", size=10)
W1_FILL = PatternFill("solid", start_color="DDEBF7")
W2_FILL = PatternFill("solid", start_color="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def style_body(ws, start_row, end_row, cols):
    for r in range(start_row, end_row + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = BORDER


# ===== Sheet 1: 커리큘럼 =====
ws1 = wb.active
ws1.title = "10일 커리큘럼"

headers = ["주차", "Day", "주제", "간략 설명", "시간 구성", "WG 대응 / 프로젝트 과제", "핵심 산출물"]
ws1.append(headers)
style_header(ws1, 1, len(headers))

rows = [
    ("W1 진입 심화", "D1", "LangChain 기초 + 고급 프롬프트 엔지니어링",
     "Chains/Tools/Structured Output 기본기 + Few-shot·CoT·Self-Consistency·Role/Format 제어·Constitutional 프롬프팅 실습",
     "이론 40% / 실습 50% / 리뷰 10%", "전 WG 공통 기반", "프롬프트 라이브러리 초안"),
    ("W1 진입 심화", "D2", "RAG 심화 패턴",
     "Hybrid(Sparse+Dense) → Router-First → Graph RAG → LLM Re-ranking → Iterative/Self-RAG를 MS Office·Confluence 데이터로 단계별 실습",
     "이론 40% / 실습 50% / 리뷰 10%", "WG-A 챗봇 / WG-C 컨플·Jira", "사내 문서 RAG POC"),
    ("W1 진입 심화", "D3", "MCP 서버 개발",
     "사내 Confluence·Jira·파일시스템을 LLM이 직접 호출하도록 MCP 서버를 직접 구현 + 권한·스키마·에러 핸들링 설계",
     "이론 40% / 실습 50% / 리뷰 10%", "WG-D 사양 자동화 (#13)", "동작하는 사내 MCP 서버"),
    ("W1 진입 심화", "D4", "LangGraph Orchestration + HITL",
     "상태 그래프·분기·재시도·사람 승인 게이트를 갖춘 멀티스텝 에이전트를 LangGraph로 구현",
     "이론 40% / 실습 50% / 리뷰 10%", "WG-A 챗봇 / WG-C Jira 자동화", "HITL 포함 에이전트 템플릿"),
    ("W1 진입 심화", "D5", "Memory Architecture & Governance",
     "Vector+Graph+KV 하이브리드 메모리 + TTL/Provenance/Audit Log + PII·역할권한으로 감사 가능한 기억 구축",
     "이론 40% / 실습 50% / 리뷰 10%", "WG-A 대화이력 / WG-C 이슈이력", "Governed Memory 레이어"),
    ("W2 프로젝트 심화", "D6", "Agentic Evals & 관측성 + 프로젝트 킥오프",
     "Promptfoo·LangSmith + 평가 데이터셋 구축 + 비용·레이턴시·품질 관측. 평가셋부터 짜고 프로젝트 시작",
     "오전 강의 / 오후 팀 프로젝트", "팀별 프로젝트 스코프·성공기준·평가셋 정의", "팀별 프로젝트 정의서 + 평가셋"),
    ("W2 프로젝트 심화", "D7", "TI 연계 심화",
     "CVE/KEV/EPSS/MITRE ATT&CK 구조 + UN R155/R156 + Auto-ISAC를 MCP로 연결해 Agentic 상관분석",
     "오전 강의 / 오후 팀 프로젝트", "팀A(보안 챗봇에 TI 주입) / 팀D(사양서에 규제·CVE 연결) 핵심 / 팀B·C 진행", "TI MCP 서버 + 연결된 에이전트"),
    ("W2 프로젝트 심화", "D8", "Red Teaming 기초 + LLM 보안",
     "PyRIT·Garak·Promptfoo 자동 Red Team 스위트 + Prompt Injection·Jailbreak + OWASP LLM Top 10 (2026)",
     "오전 강의 / 오후 팀 프로젝트", "각 팀 프로젝트에 Red Team 스위트 적용 + 결함 리포트", "Red Team 리포트 + 완화 조치"),
    ("W2 프로젝트 심화", "D9", "Adaptive / Self-Reflective Red Teaming + 고도화",
     "공격 실패→원인 요약→전략 재구성→재시도 루프 + HITL Gate 심화 + Memory Loop 결합한 자가 개선 에이전트",
     "오전 강의 / 오후 팀 프로젝트", "프로젝트 최종 고도화 + 시연 리허설", "고도화된 에이전트 v1.0"),
    ("W2 프로젝트 심화", "D10", "캡스톤 발표 + 운영 이관",
     "팀별 데모·코드리뷰·Q&A + 사내 배포·모니터링·운영 가이드(관측성·알람·롤백)",
     "오전 발표 / 오후 피어리뷰·이관", "발표 / 팀간 피어리뷰 / 사내 이관 체크리스트", "데모 + 운영 이관 README"),
]
for r in rows:
    ws1.append(r)

style_body(ws1, 2, len(rows) + 1, len(headers))
for r in range(2, 7):
    ws1.cell(row=r, column=1).fill = W1_FILL
for r in range(7, 12):
    ws1.cell(row=r, column=1).fill = W2_FILL

widths = [14, 6, 32, 60, 22, 38, 28]
for i, w in enumerate(widths, 1):
    ws1.column_dimensions[get_column_letter(i)].width = w
for r in range(2, len(rows) + 2):
    ws1.row_dimensions[r].height = 78
ws1.row_dimensions[1].height = 28
ws1.freeze_panes = "A2"

# ===== Sheet 2: WG 매핑 =====
ws2 = wb.create_sheet("HKMC WG 매핑")
wg_headers = ["WG", "주제", "필요 기술", "관련 과제"]
ws2.append(wg_headers)
style_header(ws2, 1, len(wg_headers))

wg_rows = [
    ("A", "Chatbot", "RAG, DB, 전처리", "#1 챗봇, #3, #5 챗봇, #6"),
    ("B", "검증 자동화(검토/정합성)", "프롬프트엔지니어링, LLM", "#2 사양 확인, #4 서류 리뷰, #12 문서·Jira 점검 체크"),
    ("C", "검증워킹그룹 특화", "RAG, DB, 프롬프트엔지니어링, LLM", "#8 컨플/Jira 기반 T, #9 컨플/Jira 평가, #10 JIRA 이슈 패턴 분석, #11 JIRA 이슈 관리 자동화"),
    ("D", "사양 자동화", "데이터 전처리, MCP", "#13 기능안전 요구사양서 개발"),
    ("E", "Drawing 자동화", "LLM", "#7 서류 기반 자동 drawing"),
]
for r in wg_rows:
    ws2.append(r)
style_body(ws2, 2, len(wg_rows) + 1, len(wg_headers))
for i, w in enumerate([8, 26, 30, 60], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
for r in range(2, len(wg_rows) + 2):
    ws2.row_dimensions[r].height = 40
ws2.row_dimensions[1].height = 28
ws2.freeze_panes = "A2"

# ===== Sheet 3: 팀별 심화 주제 활용 =====
ws3 = wb.create_sheet("팀별 심화 활용")
tm_headers = ["팀", "담당 WG / 과제", "Memory/Governance", "TI 연계", "Red Team"]
ws3.append(tm_headers)
style_header(ws3, 1, len(tm_headers))

tm_rows = [
    ("팀A", "WG-A Chatbot (#1,5)", "◎ 대화이력·감사", "◎ 보안 QA에 TI", "◎ Jailbreak 방어"),
    ("팀B", "WG-B 검증 자동화 (#2,4,12)", "○ 검증이력", "△", "◎ 검증봇 신뢰성"),
    ("팀C", "WG-C Jira/컨플 (#10,11)", "◎ 이슈이력·패턴", "△", "○ 데이터 누출"),
    ("팀D", "WG-D 사양 자동화 (#13)", "○ 버전 이력", "◎ R155/156 연계", "○ 사양 오류 유도"),
]
for r in tm_rows:
    ws3.append(r)
style_body(ws3, 2, len(tm_rows) + 1, len(tm_headers))
for i, w in enumerate([8, 32, 22, 22, 22], 1):
    ws3.column_dimensions[get_column_letter(i)].width = w
for r in range(2, len(tm_rows) + 2):
    ws3.row_dimensions[r].height = 32
ws3.row_dimensions[1].height = 28
ws3.freeze_panes = "A2"

note = ws3.cell(row=len(tm_rows) + 3, column=1, value="◎ 강함  ○ 중간  △ 약함")
note.font = Font(name="맑은 고딕", size=9, italic=True, color="595959")

# ===== Sheet 4: 최종 산출물 =====
ws4 = wb.create_sheet("팀 최종 산출물")
ws4.append(["#", "산출물", "설명"])
style_header(ws4, 1, 3)
out_rows = [
    (1, "작동하는 데모", "사내 환경 가정, 팀별 WG 과제 해결"),
    (2, "평가 리포트", "평가셋 + 정확도·커버리지·비용 측정 결과"),
    (3, "Red Team 리포트", "취약성 스위트 결과 + 완화 조치"),
    (4, "거버넌스 설계서", "HITL·권한·Audit Log 설계 문서"),
    (5, "운영 이관 README", "MCP 설정, 모니터링 포인트, 알려진 한계"),
]
for r in out_rows:
    ws4.append(r)
style_body(ws4, 2, len(out_rows) + 1, 3)
for i, w in enumerate([6, 26, 60], 1):
    ws4.column_dimensions[get_column_letter(i)].width = w
for r in range(2, len(out_rows) + 2):
    ws4.row_dimensions[r].height = 32
ws4.row_dimensions[1].height = 28

out_path = "C:/Users/Admin/Downloads/ai-education-proposal/projects/현대자동차/H2_심화교육_10일_커리큘럼.xlsx"
wb.save(out_path)
print(out_path)
