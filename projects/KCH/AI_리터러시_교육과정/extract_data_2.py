import sys
try:
    import pandas as pd
    from docx import Document
except ImportError:
    sys.exit(1)

out = ""
out += "=== DOCX CONTENT ===\n"
try:
    doc = Document(r"c:\Users\Admin\Downloads\ai-education-proposal\projects\KCH\AI_리터러시_교육과정\[KCH] AI 리터러시 진단 문항(최종) (2).docx")
    for para in doc.paragraphs:
        if para.text.strip():
            out += para.text[:300] + "\n"
    
    for table in doc.tables:
        for row in table.rows:
            row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            out += " | ".join(row_data)[:300] + "\n"
except Exception as e:
    out += f"Error reading docx: {e}\n"

out += "\n=== XLSX CONTENT ===\n"
try:
    excel_file = r"c:\Users\Admin\Downloads\ai-education-proposal\projects\KCH\AI_리터러시_교육과정\[KCH} AI 리터러시 진단 테스트 결과 (1).xlsx"
    xls = pd.ExcelFile(excel_file)
    out += f"Sheet names: {xls.sheet_names}\n"
    for sheet in xls.sheet_names:
        out += f"\n[Sheet: {sheet}]\n"
        df = pd.read_excel(excel_file, sheet_name=sheet)
        out += "Columns: " + str(list(df.columns)) + "\n"
        out += "Data sample (first 3 rows):\n"
        out += df.head(3).to_string() + "\n"
        
        # calculate mean scores
        numeric_cols = df.select_dtypes(include=['number']).columns
        if len(numeric_cols) > 0:
            out += "\nMean values for numeric columns:\n"
            out += df[numeric_cols].mean().to_string() + "\n"
except Exception as e:
    out += f"Error reading xlsx: {e}\n"

with open("extraction_output_utf8.txt", "w", encoding="utf-8") as f:
    f.write(out)
