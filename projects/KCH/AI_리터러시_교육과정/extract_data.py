import sys

try:
    import pandas as pd
    from docx import Document
except ImportError:
    print("Missing pandas or python-docx. Please install them.")
    sys.exit(1)

print("=== DOCX CONTENT ===")
try:
    doc = Document(r"c:\Users\Admin\Downloads\ai-education-proposal\projects\KCH\AI_리터러시_교육과정\[KCH] AI 리터러시 진단 문항(최종) (2).docx")
    for para in doc.paragraphs:
        if para.text.strip():
            print(para.text[:200]) # truncated for brevity if long
    
    for table in doc.tables:
        for row in table.rows:
            row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            print(" | ".join(row_data)[:200])
except Exception as e:
    print(f"Error reading docx: {e}")

print("\n=== XLSX CONTENT ===")
try:
    excel_file = r"c:\Users\Admin\Downloads\ai-education-proposal\projects\KCH\AI_리터러시_교육과정\[KCH} AI 리터러시 진단 테스트 결과 (1).xlsx"
    # read all sheets
    xls = pd.ExcelFile(excel_file)
    print(f"Sheet names: {xls.sheet_names}")
    for sheet in xls.sheet_names:
        print(f"\n[Sheet: {sheet}]")
        df = pd.read_excel(excel_file, sheet_name=sheet)
        print("Columns:", list(df.columns))
        print("Data sample (first 5 rows):")
        print(df.head(5).to_string())
        print("\nDescriptive Summary:")
        print(df.describe().to_string())
except Exception as e:
    print(f"Error reading xlsx: {e}")
